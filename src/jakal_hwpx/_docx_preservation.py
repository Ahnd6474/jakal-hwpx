"""Keep the original OPC package and patch supported edits without rebuilding it."""
from __future__ import annotations

from copy import deepcopy
from dataclasses import asdict
from difflib import SequenceMatcher
import io
import posixpath
from pathlib import Path
from typing import Any
import zipfile
from urllib.parse import unquote

from lxml import etree

from ._hancom_models import Hyperlink, Paragraph, Table

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _path(element: Any) -> str:
    return element.getroottree().getelementpath(element)


def _visible_slots(paragraph: Any):
    """Only the runs exposed by the common model, leaving revisions/SDTs intact."""
    from docx.text.hyperlink import Hyperlink as WordHyperlink

    slots = []
    position = 0
    for content in paragraph.iter_inner_content():
        runs = content.runs if isinstance(content, WordHyperlink) else [content]
        for run in runs:
            for node in run._r:
                name = etree.QName(node).localname
                value = node.text or "" if name == "t" else "\t" if name == "tab" else "\n" if name in {"br", "cr"} and node.get(f"{{{W}}}type", "textWrapping") == "textWrapping" else ""
                if value:
                    slots.append((_path(node), position, position + len(value), value))
                    position += len(value)
    return slots


def _replace_slot(node: etree._Element, value: str) -> None:
    parent = node.getparent()
    index = parent.index(node)
    tail = node.tail
    parent.remove(node)
    parts = value.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    nodes = []
    for line_index, line in enumerate(parts):
        if line_index:
            nodes.append(etree.Element(f"{{{W}}}br"))
        for tab_index, chunk in enumerate(line.split("\t")):
            if tab_index:
                nodes.append(etree.Element(f"{{{W}}}tab"))
            if chunk:
                text = etree.Element(f"{{{W}}}t")
                text.text = chunk
                if chunk[:1].isspace() or chunk[-1:].isspace():
                    text.set(XML_SPACE, "preserve")
                nodes.append(text)
    for offset, child in enumerate(nodes):
        parent.insert(index + offset, child)
    if nodes:
        nodes[-1].tail = tail


def _patch_paragraph(root: etree._Element, source: Any, old: str, new: str) -> None:
    if old == new:
        return
    slots = _visible_slots(source)
    if "".join(slot[3] for slot in slots) != old:
        raise ValueError("DOCX text edit cannot be mapped without changing unsupported inline content.")
    resolved = [(root.find(path), start, end, value) for path, start, end, value in slots]
    if not slots:
        paragraph = root.find(_path(source._p))
        run = etree.SubElement(paragraph, f"{{{W}}}r")
        node = etree.SubElement(run, f"{{{W}}}t")
        _replace_slot(node, new)
        return
    # Accumulate edits per original text slot, then change XML only once. This
    # preserves every run's rPr, field, bookmark, drawing and relationship.
    changes = {id(node): value for node, _, _, value in resolved}
    for tag, start, end, new_start, new_end in reversed(SequenceMatcher(None, old, new, autojunk=False).get_opcodes()):
        if tag == "equal":
            continue
        indices = [i for i, (_, a, b, _) in enumerate(resolved) if a < end and b > start]
        if not indices:
            i = next((i for i, (_, a, b, _) in enumerate(resolved) if a <= start < b), len(resolved) - 1)
            node, a, _, _ = resolved[i]
            offset = start - a
            current = changes[id(node)]
            changes[id(node)] = current[:offset] + new[new_start:new_end] + current[offset:]
        else:
            first = indices[0]
            for i in indices:
                node, a, b, _ = resolved[i]
                current = changes[id(node)]
                left, right = max(0, start - a), min(b - a, end - a)
                changes[id(node)] = current[:left] + (new[new_start:new_end] if i == first else "") + current[right:]
    for node, _, _, original in resolved:
        if changes[id(node)] != original:
            _replace_slot(node, changes[id(node)])


def replace_word_match(root: etree._Element, paragraph: Any, start: int, end: int, replacement: str) -> None:
    """Match Word's run boundary on a simple Find/Replace edit.

    Keeping a replacement in one long run can change Word's glyph placement
    compared with native Find/Replace, even when rPr is identical.
    """
    slots = _visible_slots(paragraph)
    touched = [(path, a, b) for path, a, b, _ in slots if a < end and b > start]
    runs = {root.find(path).getparent() for path, _, _ in touched}
    if len(runs) == 1:
        run = next(iter(runs))
        if all(child.tag in {f"{{{W}}}rPr", f"{{{W}}}t"} for child in run):
            run_slots = [(path, a, b, value) for path, a, b, value in slots if root.find(path).getparent() is run]
            offset = min(a for _, a, _, _ in run_slots)
            value = "".join(value for _, _, _, value in run_slots)
            parts = (value[:start - offset], replacement, value[end - offset:])
            parent, index, tail = run.getparent(), run.getparent().index(run), run.tail
            clones = []
            for part in parts:
                if not part:
                    continue
                clone = deepcopy(run)
                for node in list(clone):
                    if node.tag != f"{{{W}}}rPr":
                        clone.remove(node)
                text = etree.SubElement(clone, f"{{{W}}}t")
                _replace_slot(text, part)
                clone.tail = None
                clones.append(clone)
            parent.remove(run)
            for number, clone in enumerate(clones):
                parent.insert(index + number, clone)
            if clones:
                clones[-1].tail = tail
            return
    current = paragraph.text
    _patch_paragraph(root, paragraph, current, current[:start] + replacement + current[end:])


def _patch_run_style(run: etree._Element, style: Any, baseline: Any) -> None:
    current, previous = asdict(style), asdict(baseline)
    fields = {name for name in current if current[name] != previous[name]}
    allowed = {"bold", "italic", "height", "text_color", "shade_color", "font_names", "underline", "strikeout", "script_position"}
    if fields - allowed:
        raise ValueError("Preserving DOCX output cannot map these character-style properties: " + ", ".join(sorted(fields - allowed)))
    pr = run.find(f"{{{W}}}rPr")
    if pr is None:
        pr = etree.Element(f"{{{W}}}rPr")
        run.insert(0, pr)

    def set_property(name: str, value: str | None):
        tag = f"{{{W}}}{name}"
        for child in list(pr.findall(tag)):
            pr.remove(child)
        if value is not None:
            node = etree.SubElement(pr, tag)
            node.set(f"{{{W}}}val", value)
            return node
        return None

    for name, tag in (("bold", "b"), ("italic", "i")):
        if name in fields:
            value = current[name]
            set_property(tag, None if value is None else "1" if value else "0")
    if "height" in fields:
        height = style.height
        if height is not None and (int(height) <= 0 or int(height) % 50):
            raise ValueError("Word font size must be a positive multiple of half a point (50 HWP units).")
        set_property("sz", str(int(height) // 50) if height is not None else None)
    if "text_color" in fields:
        set_property("color", str(style.text_color).lstrip("#") if style.text_color else None)
    if "font_names" in fields:
        fonts = pr.find(f"{{{W}}}rFonts")
        if fonts is None:
            fonts = etree.SubElement(pr, f"{{{W}}}rFonts")
        for attribute, lang in (("ascii", "latin"), ("hAnsi", "latin"), ("eastAsia", "hangul")):
            if lang in style.font_names:
                fonts.set(f"{{{W}}}{attribute}", style.font_names[lang])
                fonts.attrib.pop(f"{{{W}}}{attribute}Theme", None)
    if "underline" in fields:
        underline = style.underline
        value = "none" if underline.get("type", "NONE") == "NONE" else {"SOLID": "single", "DOUBLE_SLIM": "double", "DOT": "dotted", "DASH": "dash", "WAVE": "wave"}.get(underline.get("shape", "SOLID"))
        if value is None:
            raise ValueError("Word underline mapping does not support this shape.")
        node = set_property("u", value)
        if node is not None and underline.get("color"):
            node.set(f"{{{W}}}color", underline["color"].lstrip("#"))
    if "strikeout" in fields:
        set_property("strike", "0" if style.strikeout.get("shape", "NONE") == "NONE" else "1")
    if "script_position" in fields:
        set_property("vertAlign", {"SUPERSCRIPT": "superscript", "SUBSCRIPT": "subscript", "NORMAL": "baseline"}.get(style.script_position))
    if "shade_color" in fields:
        set_property("highlight", None)
        set_property("shd", None)
        if style.shade_color and style.shade_color.lower() != "none":
            shade = etree.SubElement(pr, f"{{{W}}}shd")
            shade.set(f"{{{W}}}fill", style.shade_color.lstrip("#"))
            shade.set(f"{{{W}}}val", "clear")


class DocxSource:
    def __init__(self, target: Any, source: Any, raw: bytes) -> None:
        self.raw = raw
        self.source = source
        self.native_baseline = self._native_parts()
        self.sections = []
        self.paragraphs = []
        self.tables = []
        self.baseline = {}
        self.metadata = asdict(target.metadata)
        self.collections = {name: deepcopy(getattr(target, name)) for name in (
            "style_definitions", "paragraph_styles", "character_styles", "numbering_definitions",
            "bullet_definitions", "memo_shape_definitions",
        )}
        for section, word_section in zip(target.sections, source.sections):
            original = list(section.blocks)
            self.sections.append((section, word_section, original, deepcopy(section.settings), deepcopy(section.header_footer_blocks)))
            groups: dict[int, list[Any]] = {}
            for block in original:
                self.baseline[id(block)] = asdict(block)
                groups.setdefault(block.source_paragraph_index, []).append(block)
            for index, item in enumerate(word_section.iter_inner_content()):
                blocks = groups.get(index, [])
                if not blocks:
                    continue
                if isinstance(blocks[0], Table):
                    self.tables.append((blocks[0], item))
                else:
                    self.paragraphs.append((blocks, item))

    def _native_parts(self) -> dict[str, bytes]:
        output = io.BytesIO()
        self.source.save(output)
        with zipfile.ZipFile(output) as package:
            return {name: package.read(name) for name in package.namelist()}

    def _native_changes(self, package: zipfile.ZipFile) -> tuple[dict[str, bytes], set[str], dict[str, bytes]]:
        current = self._native_parts()
        changed = {name: data for name, data in current.items() if self.native_baseline.get(name) != data}
        deleted = set(self.native_baseline) - set(current)
        # python-docx omits unreachable parts when saving. Preserve original
        # unmodelled content-type/relationship entries while applying native edits.
        for name in list(changed):
            if name not in package.namelist() or not (name.endswith(".rels") or name == "[Content_Types].xml"):
                continue
            original = etree.fromstring(package.read(name))
            before = etree.fromstring(self.native_baseline[name])
            after = etree.fromstring(current[name])

            def key(node):
                return node.get("Id") or node.get("PartName") or node.get("Extension")

            old, new = {key(node): node for node in before}, {key(node): node for node in after}
            for record_id in set(old) | set(new):
                if record_id in old and record_id in new and etree.tostring(old[record_id]) == etree.tostring(new[record_id]):
                    continue
                for record in list(original):
                    if key(record) == record_id:
                        original.remove(record)
                if record_id in new:
                    original.append(deepcopy(new[record_id]))
            changed[name] = etree.tostring(original, encoding="UTF-8", xml_declaration=True, standalone=True)
        return changed, deleted, current

    def to_bytes(self, target: Any) -> bytes:
        if len(target.sections) != len(self.sections) or any(section is not original[0] for section, original in zip(target.sections, self.sections)):
            raise ValueError("Preserving DOCX output does not support section insertion, deletion, or reordering.")
        for name, baseline in self.collections.items():
            if name == "character_styles":
                if [s.style_id for s in getattr(target, name)] != [s.style_id for s in baseline]:
                    raise ValueError("Preserving DOCX output cannot insert or delete character styles.")
                continue
            if getattr(target, name) != baseline:
                raise ValueError(f"Preserving DOCX output cannot map edits to {name}; no file was written.")
        roots = {}
        dirty = set()
        replacements: dict[str, bytes] = {}
        package = zipfile.ZipFile(io.BytesIO(self.raw))
        native_changes, deleted, current_native = self._native_changes(package)
        replacements.update(native_changes)

        def part(name: str):
            name = name.lstrip("/")
            if name not in roots:
                data = native_changes.get(name)
                if data is None:
                    data = package.read(name) if name in package.namelist() else current_native.get(name)
                if data is None and name.endswith(".rels"):
                    data = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>'
                if data is None:
                    raise ValueError(f"DOCX package has no editable part {name!r}.")
                roots[name] = etree.fromstring(data, etree.XMLParser(resolve_entities=False, remove_blank_text=False))
            return roots[name]

        main_name = str(self.source.part.partname).lstrip("/")
        main = part(main_name)
        original_styles = {style.style_id: style for style in self.collections["character_styles"]}
        current_styles = {style.style_id: style for style in target.character_styles}
        for section, word_section, original, settings, headers in self.sections:
            if len(section.blocks) < len(original) or any(a is not b for a, b in zip(section.blocks, original)):
                raise ValueError("Preserving DOCX output supports editing original blocks and appending; block deletion/reordering is not supported.")
            if section.settings != settings or section.header_footer_blocks != headers:
                raise ValueError("Preserving DOCX output cannot map section/header/footer edits yet; no file was written.")
            for block in section.blocks[len(original):]:
                if section is not target.sections[-1] or not isinstance(block, Paragraph):
                    raise ValueError("Preserving DOCX output currently supports appending paragraphs to the last section.")
                if any(value is not None for name, value in asdict(block).items() if name != "text"):
                    raise ValueError("Preserving DOCX output cannot map formatting on newly appended paragraphs yet.")
                node = etree.Element(f"{{{W}}}p")
                run = etree.SubElement(node, f"{{{W}}}r")
                text = etree.SubElement(run, f"{{{W}}}t")
                _replace_slot(text, block.text)
                body = main.find(f"{{{W}}}body")
                final_settings = body.find(f"{{{W}}}sectPr")
                if final_settings is not None:
                    final_settings.addprevious(node)
                else:
                    body.append(node)
                dirty.add(main_name)
            for block in original:
                changed = asdict(block)
                baseline = dict(self.baseline[id(block)])
                for field in ("text",) if isinstance(block, Paragraph) else ("display_text", "target") if isinstance(block, Hyperlink) else ("cell_texts",) if isinstance(block, Table) else ():
                    changed.pop(field, None)
                    baseline.pop(field, None)
                if changed != baseline:
                    raise ValueError(f"Preserving DOCX output cannot map non-text edits to {type(block).__name__}; no file was written.")
        for blocks, paragraph in self.paragraphs:
            offset = 0
            slots = _visible_slots(paragraph)
            for block in blocks:
                original_text = self.baseline[id(block)].get("text", self.baseline[id(block)].get("display_text") or "") if isinstance(block, (Paragraph, Hyperlink)) else ""
                if isinstance(block, Paragraph) and block.char_pr_id:
                    current_style = current_styles[block.char_pr_id]
                    previous_style = original_styles[block.char_pr_id]
                    if current_style != previous_style:
                        paths = {path for path, start, end, _ in slots if start < offset + len(original_text) and end > offset}
                        for path in paths:
                            _patch_run_style(main.find(path).getparent(), current_style, previous_style)
                        if paths:
                            dirty.add(main_name)
                offset += len(original_text)
            old = "".join(self.baseline[id(b)].get("text", self.baseline[id(b)].get("display_text") or "") for b in blocks if isinstance(b, (Paragraph, Hyperlink)))
            new = "".join(b.text if isinstance(b, Paragraph) else b.display_text or "" for b in blocks if isinstance(b, (Paragraph, Hyperlink)))
            if old != new:
                _patch_paragraph(main, paragraph, old, new)
                dirty.add(main_name)
            links = iter(paragraph.hyperlinks)
            for block in blocks:
                if not isinstance(block, Hyperlink):
                    continue
                link = next(links)
                if block.target == self.baseline[id(block)]["target"]:
                    continue
                node = main.find(_path(link._hyperlink))
                if block.target.startswith("#"):
                    node.attrib.pop(f"{{{R}}}id", None)
                    node.set(f"{{{W}}}anchor", block.target[1:])
                else:
                    rel_name = str(Path(main_name).parent / "_rels" / (Path(main_name).name + ".rels")).replace("\\", "/")
                    rels = part(rel_name)
                    ids = {rel.get("Id") for rel in rels}
                    number = 1
                    while f"rIdJakal{number}" in ids:
                        number += 1
                    rid = f"rIdJakal{number}"
                    etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", Id=rid, Type=R + "/hyperlink", Target=block.target, TargetMode="External")
                    node.set(f"{{{R}}}id", rid)
                    node.attrib.pop(f"{{{W}}}anchor", None)
                    dirty.add(rel_name)
                dirty.add(main_name)
        for block, table in self.tables:
            baseline = self.baseline[id(block)]["cell_texts"]
            for row, texts in enumerate(block.cell_texts):
                for column, value in enumerate(texts):
                    if value == baseline[row][column]:
                        continue
                    cell = table.cell(row, column)
                    if len(cell.paragraphs) != 1:
                        raise ValueError("Preserving DOCX output requires paragraph-aware edits for multi-paragraph cells.")
                    _patch_paragraph(main, cell.paragraphs[0], cell.text, value)
                    dirty.add(main_name)
        if asdict(target.metadata) != self.metadata:
            metadata = asdict(target.metadata)
            core_name = "docProps/core.xml"
            core = part(core_name)
            for field, tag in {
                "title": "{http://purl.org/dc/elements/1.1/}title", "creator": "{http://purl.org/dc/elements/1.1/}creator",
                "subject": "{http://purl.org/dc/elements/1.1/}subject", "description": "{http://purl.org/dc/elements/1.1/}description",
                "language": "{http://purl.org/dc/elements/1.1/}language", "keyword": "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}keywords",
            }.items():
                if metadata[field] == self.metadata[field]:
                    continue
                node = core.find(tag)
                if node is None:
                    node = etree.SubElement(core, tag)
                node.text = metadata[field]
                dirty.add(core_name)
            for field in set(metadata) - {"title", "creator", "subject", "description", "language", "keyword"}:
                if metadata[field] != self.metadata[field]:
                    raise ValueError(f"Preserving DOCX output cannot map metadata {field!r}.")
        if not dirty and not replacements and not deleted:
            package.close()
            return self.raw
        for name in dirty:
            replacements[name] = etree.tostring(roots[name], encoding="UTF-8", xml_declaration=True, standalone=True)
        # Reading python-docx's properties can lazily materialize a missing core
        # or styles part. When it is actually edited, also add its declarations;
        # otherwise the new XML would be an unreachable orphan in the package.
        added = set(replacements) - set(package.namelist())
        declarations = set()
        native_types = etree.fromstring(current_native["[Content_Types].xml"])
        for name in added:
            types = part("[Content_Types].xml")
            extension = name.rsplit(".", 1)[-1]
            for record in native_types:
                if record.get("PartName") == "/" + name or record.get("Extension") == extension:
                    key, value = ("PartName", record.get("PartName")) if record.get("PartName") else ("Extension", record.get("Extension"))
                    if not any(node.get(key) == value for node in types):
                        types.append(deepcopy(record))
                        declarations.add("[Content_Types].xml")
            for rel_name, raw in current_native.items():
                if not rel_name.endswith(".rels"):
                    continue
                base = rel_name.split("/_rels/", 1)[0] if "/_rels/" in rel_name else ""
                for relation in etree.fromstring(raw):
                    if relation.get("TargetMode") == "External":
                        continue
                    target = unquote(relation.get("Target", "").split("#", 1)[0])
                    resolved = posixpath.normpath(target if target.startswith("/") else posixpath.join(base, target)).lstrip("/")
                    if resolved != name:
                        continue
                    rels = part(rel_name)
                    if not any(node.get("Id") == relation.get("Id") for node in rels):
                        rels.append(deepcopy(relation))
                    declarations.add(rel_name)
        for name in declarations:
            replacements[name] = etree.tostring(roots[name], encoding="UTF-8", xml_declaration=True, standalone=True)
        output = io.BytesIO()
        with package, zipfile.ZipFile(output, "w") as result:
            result.comment = package.comment
            for info in package.infolist():
                if info.filename not in deleted:
                    result.writestr(info, replacements.get(info.filename, package.read(info)))
            for name in set(replacements) - set(package.namelist()):
                result.writestr(name, replacements[name])
        return output.getvalue()


def capture_docx_source(target: Any, source: Any, path: Path) -> None:
    target._docx_source = DocxSource(target, source, path.read_bytes())
