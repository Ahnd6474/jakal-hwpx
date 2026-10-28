"""Adapt python-docx objects to the existing Hancom document model.

This module is loaded only for DOCX input, so python-docx remains optional.
"""
from __future__ import annotations

import warnings
import json
from copy import deepcopy
from pathlib import Path
from typing import TYPE_CHECKING, Any

from ._hancom_models import (
    CharacterStyle,
    Equation,
    HancomBlock,
    HancomMetadata,
    HancomSection,
    HeaderFooter,
    Hyperlink,
    Paragraph,
    ParagraphStyle,
    Picture,
    SectionSettings,
    StyleDefinition,
    Table,
)

if TYPE_CHECKING:
    from .hancom_document import HancomDocument


def populate_from_docx(target: HancomDocument, path: Path) -> None:
    try:
        from docx import Document
    except ModuleNotFoundError as error:
        if error.name != "docx":
            raise
        raise ModuleNotFoundError(
            'DOCX input requires python-docx. Install it with: pip install "jakal-hwpx[docx]"'
        ) from error

    source = Document(str(path))
    from docx.table import Table as DocxTable

    reader = _DocxReader(target)
    reader.numbering_root = source.part.numbering_part.element
    for relationship in source.part.rels.values():
        if relationship.reltype.endswith("/theme") and not relationship.is_external:
            from lxml import etree

            reader.theme = etree.fromstring(relationship.target_part.blob)
            break
    props = source.core_properties
    target.metadata = HancomMetadata(
        title=props.title or None,
        language=props.language or None,
        creator=props.author or None,
        subject=props.subject or None,
        description=props.comments or None,
        lastsaveby=props.last_modified_by or None,
        created=props.created.isoformat() if props.created else None,
        modified=props.modified.isoformat() if props.modified else None,
        keyword=props.keywords or None,
    )
    defaults = source.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")
    if defaults:
        reader.default_run_properties = defaults[0]
    defaults = source.styles.element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")
    if defaults:
        from docx.oxml import OxmlElement
        from docx.text.parfmt import ParagraphFormat

        default_paragraph = OxmlElement("w:p")
        default_paragraph.append(deepcopy(defaults[0]))
        reader.default_paragraph_format = ParagraphFormat(default_paragraph)
    reader.check_unsupported(source.element)
    inherited_headers: dict[tuple[str, str], str] = {}
    for source_section in source.sections:
        section = HancomSection(settings=_section_settings(source_section))
        target.sections.append(section)
        for paragraph_index, item in enumerate(source_section.iter_inner_content()):
            if isinstance(item, DocxTable):
                block = reader.table(item)
                _mark_source(block, paragraph_index, 0)
                section.blocks.append(block)
            else:
                # A bare section-break carrier is structural, not an extra blank
                # body paragraph. This matches Word's native import into Hancom.
                if item._p.xpath("./w:pPr/w:sectPr") and not item._p.xpath("./*[not(self::w:pPr)]"):
                    continue
                section.blocks.extend(reader.paragraph(item, paragraph_index))
        reader.headers(source_section, section, inherited_headers, source.settings)
    if not target.sections:
        target.sections.append(HancomSection())
    from ._docx_preservation import capture_docx_source

    capture_docx_source(target, source, path)


def _units(length: int) -> int:
    # 914400 EMU/inch, 7200 HWP units/inch.
    return round(int(length) / 127)


def _cell_margins(node: Any) -> dict[str, int]:
    """Read direct OOXML cell margins (twips) into HWP's 1/7200-inch units."""
    from docx.oxml.ns import qn

    margins: dict[str, int] = {}
    if node is None:
        return margins
    for side, names in (
        ("left", ("left", "start")), ("right", ("right", "end")),
        ("top", ("top",)), ("bottom", ("bottom",)),
    ):
        for name in names:
            edge = node.find(qn("w:" + name))
            if edge is None or edge.get(qn("w:type"), "dxa") != "dxa":
                continue
            width = edge.get(qn("w:w"))
            if width is not None:
                margins[side] = int(width) * 5
                break
    return margins


def _border_styles(node: Any) -> dict[str, dict[str, str | int]]:
    from docx.oxml.ns import qn

    borders: dict[str, dict[str, str | int]] = {}
    if node is None:
        return borders
    for side, names in (
        ("left", ("left", "start")), ("right", ("right", "end")),
        ("top", ("top",)), ("bottom", ("bottom",)),
    ):
        for name in names:
            edge = node.find(qn("w:" + name))
            if edge is None:
                continue
            borders[side] = {
                "style": edge.get(qn("w:val"), "single"),
                "size": int(edge.get(qn("w:sz"), "4")),
                "color": edge.get(qn("w:color"), "000000"),
            }
            break
    return borders


def _section_settings(section: Any) -> SectionSettings:
    settings = SectionSettings(
        page_width=_units(section.page_width) if section.page_width is not None else None,
        page_height=_units(section.page_height) if section.page_height is not None else None,
        landscape="LANDSCAPE" if section.orientation == 1 else "WIDELY",
    )
    for name, attribute in (
        ("left", "left_margin"), ("right", "right_margin"),
        ("top", "top_margin"), ("bottom", "bottom_margin"),
        ("header", "header_distance"), ("footer", "footer_distance"), ("gutter", "gutter"),
    ):
        value = getattr(section, attribute)
        if value is not None:
            settings.margins[name] = _units(value)
    # Word measures both the body and header/footer from the paper edge. HWP's
    # top/bottom margin is the outer margin, with a separate reserved distance
    # from the header/footer to the body. Preserve the total body offset.
    for edge, story, total, distance in (
        ("top", "header", section.top_margin, section.header_distance),
        ("bottom", "footer", section.bottom_margin, section.footer_distance),
    ):
        if total is not None and distance is not None:
            outer = min(_units(total), _units(distance))
            settings.margins[edge] = outer
            settings.margins[story] = _units(total) - outer
    return settings


def _mark_source(block: HancomBlock | HeaderFooter, index: int, order: int) -> None:
    # Use the same grouping metadata as the HWP/HWPX readers.
    setattr(block, "source_paragraph_index", index)
    setattr(block, "source_order", order)


def _style_chain(style: Any):
    seen: set[str] = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


class _DocxReader:
    def __init__(self, target: HancomDocument) -> None:
        self.target = target
        self.default_run_properties: Any = None
        self.default_paragraph_format: Any = None
        self._paragraph_styles: dict[str, str] = {}
        self._character_styles: dict[str, str] = {}
        self._named_styles: dict[str, str] = {}
        self.theme: Any = None
        self.numbering_root: Any = None
        self._list_counters: dict[tuple[str, int], int] = {}

    def warn(self, message: str) -> None:
        if message not in self.target.import_warnings:
            self.target.import_warnings.append(message)
            if not getattr(self.target, "_preserves_word_details", False):
                warnings.warn(message, UserWarning, stacklevel=3)

    def check_unsupported(self, element: Any) -> None:
        for expression, feature in (
            (".//w:ins | .//w:del | .//w:moveFrom | .//w:moveTo", "tracked revisions"),
            (".//w:sdt | .//w:customXml | .//w:altChunk", "content controls or embedded document content"),
            (".//w:footnoteReference | .//w:endnoteReference", "footnotes/endnotes"),
            (".//w:commentRangeStart | .//w:commentReference", "comments"),
            (".//w:fldSimple | .//w:fldChar", "fields (only directly exposed cached text is retained)"),
            (".//w:bookmarkStart", "bookmarks"),
            (".//w:pict | .//w:object", "legacy drawings or embedded objects"),
            ('.//w:br[@w:type="page" or @w:type="column"]', "explicit page/column breaks"),
        ):
            if element.xpath(expression):
                self.warn(f"DOCX import does not preserve {feature}.")

    def paragraph(self, source: Any, index: int) -> list[HancomBlock]:
        from docx.drawing import Drawing
        from docx.text.run import Run as DocxRun
        from docx.text.hyperlink import Hyperlink as DocxHyperlink
        from lxml import etree

        para_id = self.paragraph_style(source)
        style_id = self.named_style(source)
        blocks: list[HancomBlock] = []

        def add_text(text: str, char_id: str | None) -> None:
            if not text:
                return
            if blocks and isinstance(blocks[-1], Paragraph) and blocks[-1].char_pr_id == char_id:
                blocks[-1].text += text
            else:
                blocks.append(Paragraph(text, style_id=style_id, para_pr_id=para_id, char_pr_id=char_id))

        word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
        for node in source._p:
            tag = etree.QName(node)
            if tag.namespace == math_ns and tag.localname in {"oMath", "oMathPara"}:
                from ._equation_conversion import omml_to_hwp

                unsupported = None
                try:
                    script = omml_to_hwp(node)
                except ValueError as error:
                    unsupported = str(error)
                    script = "".join(node.xpath(".//m:t/text()", namespaces={"m": math_ns}))
                    self.warn(unsupported)
                if script or unsupported:
                    equation = Equation(script=script)
                    if unsupported:
                        setattr(equation, "_unsupported_word_math", unsupported)
                    blocks.append(equation)
                continue
            if tag.namespace != word_ns or tag.localname not in {"r", "hyperlink"}:
                continue
            item = DocxRun(node, source) if tag.localname == "r" else DocxHyperlink(node, source)
            if isinstance(item, DocxHyperlink):
                target = item.url or (f"#{item.fragment}" if item.fragment else "")
                if target:
                    hyperlink = Hyperlink(target, display_text=item.text)
                    styled = [Paragraph(run.text, char_pr_id=self.character_style(source, run)) for run in item.runs]
                    setattr(hyperlink, "_docx_hyperlink_runs", (item.text, styled))
                    blocks.append(hyperlink)
                else:
                    add_text(item.text, None)
                for run in item.runs:
                    self.character_style(source, run)
                    if run._r.xpath(".//w:drawing"):
                        self.warn("DOCX import does not preserve images inside hyperlinks.")
                continue
            char_id = self.character_style(source, item)
            for content in item.iter_inner_content():
                if isinstance(content, str):
                    add_text(content, char_id)
                elif isinstance(content, Drawing):
                    picture = self.picture(content)
                    if picture is not None:
                        blocks.append(picture)
        if not blocks:
            from docx.oxml import OxmlElement

            char_id = self.character_style(source, DocxRun(OxmlElement("w:r"), source))
            blocks.append(Paragraph("", style_id=style_id, para_pr_id=para_id, char_pr_id=char_id))
        # A leading empty paragraph carries alignment for picture/link-only paragraphs.
        elif not isinstance(blocks[0], Paragraph):
            blocks.insert(0, Paragraph("", style_id=style_id, para_pr_id=para_id))
        marker = self._numbering_marker(source)
        if marker:
            first_paragraph = next((block for block in blocks if isinstance(block, Paragraph)), None)
            if first_paragraph is not None:
                first_paragraph.list_marker = marker
            self.warn("DOCX list labels are materialized as static text in HWP/HWPX; automatic numbering is not retained.")
        for order, block in enumerate(blocks):
            _mark_source(block, index, order)
        return blocks

    def _numbering_marker(self, source: Any) -> str | None:
        from docx.oxml.ns import qn

        numbering = self.numbering_root
        if numbering is None:
            return None
        properties = [source._p, *(style.element for style in _style_chain(source.style))]
        number_property = next(
            (node for element in properties if (node := element.find(f"./{qn('w:pPr')}/{qn('w:numPr')}")) is not None),
            None,
        )
        if number_property is None:
            return None
        number_id_node = number_property.find(qn("w:numId"))
        if number_id_node is None:
            return None
        number_id = number_id_node.get(qn("w:val"), "0")
        if number_id == "0":
            return None
        level_node = number_property.find(qn("w:ilvl"))
        level = int(level_node.get(qn("w:val"), "0")) if level_node is not None else 0
        number = next(
            (node for node in numbering if node.tag == qn("w:num") and node.get(qn("w:numId")) == number_id),
            None,
        )
        if number is None:
            return None
        abstract_ref = number.find(qn("w:abstractNumId"))
        if abstract_ref is None:
            return None
        abstract_id = abstract_ref.get(qn("w:val"))
        abstract = next(
            (node for node in numbering if node.tag == qn("w:abstractNum") and node.get(qn("w:abstractNumId")) == abstract_id),
            None,
        )
        if abstract is None:
            return None
        level_definition = next(
            (node for node in abstract if node.tag == qn("w:lvl") and node.get(qn("w:ilvl")) == str(level)),
            None,
        )
        if level_definition is None:
            return None
        format_node = level_definition.find(qn("w:numFmt"))
        text_node = level_definition.find(qn("w:lvlText"))
        if format_node is None or text_node is None:
            return None
        number_format = format_node.get(qn("w:val"), "")
        marker = text_node.get(qn("w:val"), "")
        if number_format == "bullet":
            return marker.replace("\uf0b7", "•").replace("\uf0a7", "▪") or "•"
        if number_format == "decimal":
            key = (number_id, level)
            start_node = level_definition.find(qn("w:start"))
            start = int(start_node.get(qn("w:val"), "1")) if start_node is not None else 1
            value = self._list_counters.get(key, start - 1) + 1
            self._list_counters[key] = value
            for counter_key in list(self._list_counters):
                if counter_key[0] == number_id and counter_key[1] > level:
                    del self._list_counters[counter_key]
            for parent_level in range(level + 1):
                marker = marker.replace(f"%{parent_level + 1}", str(self._list_counters.get((number_id, parent_level), 1)))
            return marker
        self.warn(f"DOCX list format {number_format!r} is not yet materialized in HWP/HWPX.")
        return None

    def paragraph_style(self, source: Any, *, include_direct: bool = True) -> str:
        formats = [style.paragraph_format for style in _style_chain(source.style)]
        if include_direct:
            formats.insert(0, source.paragraph_format)
        if self.default_paragraph_format is not None:
            formats.append(self.default_paragraph_format)

        def effective(name: str):
            return next((getattr(fmt, name) for fmt in formats if getattr(fmt, name) is not None), None)

        alignment = effective("alignment")
        alignment_name = {0: "LEFT", 1: "CENTER", 2: "RIGHT", 3: "JUSTIFY", 4: "DISTRIBUTE"}.get(alignment, "LEFT")
        if alignment is not None and alignment_name is None:
            self.warn("DOCX import does not preserve this paragraph alignment; using the default.")
        style = ParagraphStyle(alignment_horizontal=alignment_name)
        for attribute, key in (
            ("left_indent", "left"), ("right_indent", "right"), ("first_line_indent", "intent"),
            ("space_before", "prev"), ("space_after", "next"),
        ):
            value = effective(attribute)
            if value is not None:
                style.margins[key] = str(_units(value))
        spacing = effective("line_spacing")
        if isinstance(spacing, float):
            style.line_spacing = round(spacing * 100)
            style.line_spacing_type = "PERCENT"
        elif spacing is not None:
            style.line_spacing = _units(spacing)
            style.line_spacing_type = "AT_LEAST" if effective("line_spacing_rule") == 3 else "FIXED"
        for attribute, key in (
            ("keep_with_next", "keepWithNext"), ("keep_together", "keepLines"),
            ("page_break_before", "pageBreakBefore"), ("widow_control", "widowOrphan"),
        ):
            value = effective(attribute)
            if value is not None:
                style.break_setting[key] = "1" if value else "0"
        tabs: Any = next((fmt.tab_stops for fmt in formats if len(fmt.tab_stops)), [])
        tab_stops = [{"pos": _units(tab.position), "type": {0: "LEFT", 1: "CENTER", 2: "RIGHT", 3: "DECIMAL", 4: "BAR"}.get(tab.alignment, "LEFT"),
                      "leader": {0: "NONE", 1: "DOT", 2: "DASH", 3: "SOLID", 4: "DOUBLE_SLIM"}.get(tab.leader, "NONE")} for tab in tabs]
        setattr(style, "_docx_tab_stops", tab_stops)
        key = repr(style) + repr(tab_stops)
        if key not in self._paragraph_styles:
            style.style_id = str(len(self._paragraph_styles) + 1)
            self.target.paragraph_styles.append(style)
            self._paragraph_styles[key] = style.style_id
        return self._paragraph_styles[key]

    def named_style(self, source: Any) -> str | None:
        style = source.style
        if style is None:
            return None
        if style.style_id not in self._named_styles:
            from docx.oxml import OxmlElement
            from docx.text.run import Run

            style_id = str(len(self._named_styles) + 1)
            self._named_styles[style.style_id] = style_id
            self.target.style_definitions.append(StyleDefinition(
                style.name, style_id=style_id, para_pr_id=self.paragraph_style(source, include_direct=False),
                char_pr_id=self.character_style(source, Run(OxmlElement("w:r"), source)),
            ))
        return self._named_styles[style.style_id]

    def character_style(self, paragraph: Any, run: Any) -> str | None:
        from docx.oxml.ns import qn

        properties = [run._r.rPr]
        properties.extend(style.element.rPr for style in _style_chain(run.style))
        properties.extend(style.element.rPr for style in _style_chain(paragraph.style))
        properties.append(self.default_run_properties)
        properties = [prop for prop in properties if prop is not None]

        def value(tag: str) -> str | None:
            for prop in properties:
                element = prop.find(qn(f"w:{tag}"))
                if element is not None:
                    return element.get(qn("w:val"))
            return None

        color = value("color")
        color = f"#{color.upper()}" if color and len(color) == 6 and all(c in "0123456789abcdefABCDEF" for c in color) else None
        size = value("sz")
        height = int(size) * 50 if size is not None else None
        def enabled(tag: str) -> bool | None:
            for prop in properties:
                node = prop.find(qn(f"w:{tag}"))
                if node is not None:
                    return node.get(qn("w:val"), "1").lower() not in {"0", "false", "off", "none"}
            return None

        fonts = {}
        for lang, attribute in (("latin", "ascii"), ("hangul", "eastAsia")):
            for prop in properties:
                node = prop.find(qn("w:rFonts"))
                if node is None:
                    continue
                face = node.get(qn("w:" + attribute))
                theme = node.get(qn("w:" + attribute + "Theme"))
                if not face and theme and self.theme is not None:
                    kind = "majorFont" if theme.startswith("major") else "minorFont"
                    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                    font = "a:latin" if lang == "latin" else 'a:font[@script="Hang"]'
                    names = self.theme.xpath(f".//a:fontScheme/a:{kind}/{font}/@typeface", namespaces=ns)
                    face = names[0] if names else None
                if face:
                    fonts[lang] = face
                    break
        for lang in ("hanja", "japanese"):
            if "hangul" in fonts:
                fonts[lang] = fonts["hangul"]
        for lang in ("other", "symbol", "user"):
            if "latin" in fonts:
                fonts[lang] = fonts["latin"]
        underline = value("u")
        underline_shapes = {"single": "SOLID", "double": "DOUBLE_SLIM", "dotted": "DOT", "dash": "DASH", "wave": "WAVE"}
        if underline not in (None, "none", *underline_shapes):
            self.warn(f"DOCX underline style {underline!r} is approximated in model conversion.")
        highlight = value("highlight")
        text_scale = value("w")
        font_ratios = {}
        if text_scale is not None:
            scale = int(text_scale)
            if 50 <= scale <= 200:
                font_ratios = {lang: scale for lang in ("hangul", "latin", "hanja", "japanese", "other", "symbol", "user")}
            else:
                self.warn("DOCX character width outside the HWP 50–200% range is not preserved.")
        shades = {"yellow": "#FFFF00", "green": "#00FF00", "cyan": "#00FFFF", "magenta": "#FF00FF", "blue": "#0000FF", "red": "#FF0000", "darkBlue": "#000080", "darkCyan": "#008080", "darkGreen": "#008000", "darkMagenta": "#800080", "darkRed": "#800000", "darkYellow": "#808000", "darkGray": "#808080", "lightGray": "#C0C0C0", "black": "#000000", "white": "#FFFFFF", "none": "none"}
        style = CharacterStyle(
            text_color=color, height=height, font_names=fonts, font_ratios=font_ratios,
            bold=enabled("b"), italic=enabled("i"),
            underline={"type": "BOTTOM" if underline not in (None, "none") else "NONE", "shape": underline_shapes.get(underline or "", "SOLID"), "color": color or "#000000"},
            strikeout={"shape": "SOLID" if enabled("strike") or enabled("dstrike") else "NONE", "color": color or "#000000"},
            shade_color=shades.get(highlight) if highlight else None,
            script_position={"superscript": "SUPERSCRIPT", "subscript": "SUBSCRIPT"}.get(value("vertAlign") or "", "NORMAL"),
        )
        if any(prop.xpath("./w:spacing | ./w:position | ./w:shd | ./w:color[@w:themeColor]") for prop in properties):
            self.warn("DOCX conversion may approximate character spacing, positioning, shading, or theme colors.")
        key = json.dumps({"color": color, "height": height, "fonts": fonts, "ratios": font_ratios, "bold": style.bold, "italic": style.italic, "underline": style.underline, "strikeout": style.strikeout, "shade": style.shade_color, "position": style.script_position}, sort_keys=True)
        if key not in self._character_styles:
            style_id = str(len(self._character_styles) + 1)
            style.style_id = style_id
            self.target.character_styles.append(style)
            self._character_styles[key] = style_id
        return self._character_styles[key]

    def picture(self, drawing: Any) -> Picture | None:
        image_ids = drawing._drawing.xpath(".//pic:blipFill/a:blip/@r:embed")
        if not image_ids:
            self.warn("DOCX import does not preserve linked images, charts, SmartArt, or other non-picture drawings.")
            return None
        image_part = drawing.part.related_parts[image_ids[0]]
        extents = drawing._drawing.xpath("./wp:inline/wp:extent | ./wp:anchor/wp:extent")
        width = _units(int(extents[0].get("cx"))) if extents else 7200
        height = _units(int(extents[0].get("cy"))) if extents else 7200
        if drawing._drawing.xpath("./wp:anchor | .//a:srcRect | .//a:xfrm[@rot]"):
            self.warn("DOCX import retains picture data and size; floating placement, crop, and rotation are not preserved.")
        descriptions = drawing._drawing.xpath("./wp:inline/wp:docPr/@descr | ./wp:anchor/wp:docPr/@descr")
        return Picture(
            name=Path(str(image_part.partname)).name, data=image_part.blob,
            extension=Path(str(image_part.partname)).suffix.lstrip("."),
            width=width, height=height, shape_comment=descriptions[0] if descriptions else None,
            layout={"treatAsChar": "1"},
        )

    def table(self, source: Any) -> Table:
        from docx.oxml.ns import qn

        rows, cols = len(source.rows), len(source.columns)
        if not rows or not cols:
            raise ValueError("DOCX contains a table with no rows or columns.")
        if source._tbl.xpath("./w:tblPr/w:tblStyle"):
            self.warn("DOCX table styles are not fully preserved; direct cell borders and shading are retained.")
        texts = [["" for _ in range(cols)] for _ in range(rows)]
        coordinates: dict[Any, list[tuple[int, int]]] = {}
        vertical_aligns: dict[tuple[int, int], str] = {}
        cell_margins: dict[tuple[int, int], dict[str, int]] = {}
        cell_fill_colors: dict[tuple[int, int], str] = {}
        cell_border_styles: dict[tuple[int, int], dict[str, dict[str, str | int]]] = {}
        table_margins = {}
        table_borders = {}
        for style in reversed(list(_style_chain(source.style))):
            properties = style.element.find(qn("w:tblPr"))
            if properties is not None:
                table_margins.update(_cell_margins(properties.find(qn("w:tblCellMar"))))
                table_borders.update(_border_styles(properties.find(qn("w:tblBorders"))))
        table_margins.update(_cell_margins(source._tbl.tblPr.find(qn("w:tblCellMar"))))
        table_borders.update(_border_styles(source._tbl.tblPr.find(qn("w:tblBorders"))))
        rich_cells = {}
        for row_index, row in enumerate(source.rows):
            for column, cell in enumerate(row.cells, start=row.grid_cols_before):
                if column >= cols:
                    raise ValueError("DOCX table cells exceed the declared column grid.")
                positions = coordinates.setdefault(cell._tc, [])
                if not positions:
                    texts[row_index][column] = cell.text
                    rich_cells[(row_index, column)] = (cell.text, self.container_blocks(cell))
                    align = {0: "TOP", 1: "CENTER", 3: "BOTTOM"}.get(cell.vertical_alignment)
                    vertical_aligns[(row_index, column)] = align or "TOP"
                    properties = cell._tc.tcPr
                    overrides = _cell_margins(properties.find(qn("w:tcMar")) if properties is not None else None)
                    if overrides:
                        cell_margins[(row_index, column)] = {**table_margins, **overrides} if table_margins else overrides
                    borders = _border_styles(properties.find(qn("w:tcBorders")) if properties is not None else None)
                    if table_borders or borders:
                        cell_border_styles[(row_index, column)] = {**table_borders, **borders}
                        if any(
                            spec.get("style") not in {"single", "dashed", "dotted", "double", "none", "nil"}
                            for spec in cell_border_styles[(row_index, column)].values()
                        ):
                            self.warn("DOCX table border style is approximated as a solid line.")
                    shading = properties.find(qn("w:shd")) if properties is not None else None
                    if shading is not None:
                        color = shading.get(qn("w:fill"), "")
                        if len(color) == 6 and all(char in "0123456789abcdefABCDEF" for char in color):
                            cell_fill_colors[(row_index, column)] = f"#{color.upper()}"
                        elif color not in ("", "auto") or shading.get(qn("w:themeFill")):
                            self.warn("DOCX theme or patterned table cell shading is not preserved.")
                positions.append((row_index, column))
        spans = {}
        for positions in coordinates.values():
            first_row, first_col = positions[0]
            row_span = max(row for row, _ in positions) - first_row + 1
            col_span = max(col for _, col in positions) - first_col + 1
            if row_span > 1 or col_span > 1:
                spans[(first_row, first_col)] = (row_span, col_span)
        widths = [column.width for column in source.columns]
        heights = [row.height for row in source.rows]
        result = Table(
            rows=rows, cols=cols, cell_texts=texts,
            col_widths=[_units(width) for width in widths] if all(w is not None for w in widths) else None,
            row_heights=[_units(height) if height is not None else 0 for height in heights] if any(h is not None for h in heights) else None,
            cell_spans=spans or None, cell_vertical_aligns=vertical_aligns or None,
            table_margins=table_margins or None, cell_margins=cell_margins or None,
            cell_fill_colors=cell_fill_colors or None, cell_border_styles=cell_border_styles or None,
            repeat_header=bool(source.rows[0]._tr.xpath("./w:trPr/w:tblHeader[not(@w:val='0' or @w:val='false' or @w:val='off')]")),
        )
        setattr(result, "_docx_cell_content", rich_cells)
        return result

    def container_blocks(self, container: Any) -> list[HancomBlock]:
        from docx.table import Table as WordTable

        blocks: list[HancomBlock] = []
        for index, item in enumerate(container.iter_inner_content()):
            values = [self.table(item)] if isinstance(item, WordTable) else self.paragraph(item, index)
            for order, block in enumerate(values):
                setattr(block, "nested_paragraph_index", index)
                setattr(block, "nested_source_order", order)
            blocks.extend(values)
        return blocks

    def headers(self, source: Any, section: HancomSection, inherited: dict[tuple[str, str], str], settings: Any) -> None:
        if source.different_first_page_header_footer:
            self.warn("DOCX import does not preserve separate first-page headers/footers.")
        for kind in ("header", "footer"):
            for variant, prefix in (("default", ""), ("even", "even_page_"), ("first", "first_page_")):
                story = getattr(source, prefix + kind)
                key = (kind, variant)
                if story.is_linked_to_previous and key not in inherited:
                    continue  # Reading absent stories must not create Word parts.
                if not story.is_linked_to_previous:
                    self.check_unsupported(story._element)
                    inherited[key] = "\n".join(paragraph.text for paragraph in story.paragraphs)
                text = inherited.get(key, "")
                active = variant == "default" or (
                    settings.odd_and_even_pages_header_footer if variant == "even" else source.different_first_page_header_footer
                )
                if not active or (not text and not story.tables and not story._element.xpath(".//w:drawing")):
                    continue
                if variant == "first":
                    self.warn("DOCX import does not preserve separate first-page headers/footers.")
                    continue
                page_type = "EVEN" if variant == "even" else ("ODD" if settings.odd_and_even_pages_header_footer else "BOTH")
                header = HeaderFooter(kind, text, apply_page_type=page_type)
                header.nested_blocks = list(self.container_blocks(story))
                _mark_source(header, 0, -10 + len(section.header_footer_blocks))
                section.header_footer_blocks.append(header)
