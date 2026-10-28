"""Build a deterministic DOCX fixture for DOCX -> HWPX/HWP fidelity work."""
from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "examples" / "samples" / "docx" / "representative.docx"
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII="
)


def _add_cell_shading(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def _add_cell_border(cell, edge: str, color: str, size: int = 12) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:color"), color)
    borders.append(border)


def _add_math(paragraph) -> None:
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x+y=3"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)


def build(path: Path = OUTPUT) -> Path:
    document = Document()
    document.core_properties.title = "DOCX fidelity benchmark"
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = "Standard header"
    section.first_page_header.paragraphs[0].text = "First-page header"
    section.footer.paragraphs[0].text = "Benchmark footer"

    title = document.add_heading("DOCX fidelity benchmark", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph = document.add_paragraph()
    normal = paragraph.add_run("Normal, ")
    normal.font.name = "Arial"
    bold = paragraph.add_run("bold red")
    bold.bold = True
    bold.font.color.rgb = RGBColor(200, 0, 0)
    bold.font.size = Pt(14)
    italic = paragraph.add_run(", italic blue")
    italic.italic = True
    italic.font.color.rgb = RGBColor(0, 70, 180)
    paragraph.add_run(", and underlined.").underline = True
    paragraph.paragraph_format.space_after = Pt(12)

    document.add_paragraph("Bullet one", style="List Bullet")
    document.add_paragraph("Bullet two", style="List Bullet")
    document.add_paragraph("Number one", style="List Number")
    document.add_paragraph("Number two", style="List Number")

    table = document.add_table(rows=3, cols=2)
    table.columns[0].width = Inches(2.3)
    table.columns[1].width = Inches(3.9)
    for row in table.rows:
        row.height = Inches(0.34)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
    table.cell(1, 0).text = "Shaded label"
    table.cell(1, 1).text = "Value"
    table.cell(2, 0).text = "Bordered"
    table.cell(2, 1).text = "42"
    table.cell(1, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _add_cell_shading(table.cell(0, 0), "D9EAF7")
    _add_cell_shading(table.cell(1, 0), "FFF2CC")
    for cell in (table.cell(0, 0), table.cell(1, 0), table.cell(1, 1), table.cell(2, 0), table.cell(2, 1)):
        for edge in ("top", "left", "bottom", "right"):
            _add_cell_border(cell, edge, "4472C4", 8)

    image_paragraph = document.add_paragraph("Inline image: ")
    image_paragraph.add_run().add_picture(BytesIO(PNG), width=Inches(0.3), height=Inches(0.3))
    math_paragraph = document.add_paragraph("Equation: ")
    _add_math(math_paragraph)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


if __name__ == "__main__":
    print(build())
