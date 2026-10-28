"""Reference marks keep their own formatting across Word and HWPX boundaries."""
from __future__ import annotations

from pathlib import Path

from docx import Document as WordFile
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.parts.story import StoryPart
from docx.shared import Pt, RGBColor
from docx.text.run import Run
from lxml import etree
import pytest

from jakal_hwpx import Document, HwpxDocument
from jakal_hwpx.hancom_document import Note
from jakal_hwpx.namespaces import NS


def _script(font, value):
    if value is not None:
        if value == "baseline":
            node = OxmlElement("w:vertAlign")
            node.set(qn("w:val"), "baseline")
            font._element.get_or_add_rPr().append(node)
        elif value == "subscript":
            font.subscript = True
        else:
            font.superscript = value == "superscript"


def _source(kind, body_script, marker_script):
    word = WordFile()
    style = word.styles.add_style("Body Reference", WD_STYLE_TYPE.CHARACTER)
    style.font.name = "Arial"
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor.from_string("224466")
    style.font.italic = True
    _script(style.font, body_script)
    paragraph = word.add_paragraph("Before ")
    reference = paragraph.add_run()
    reference.style = style
    node = OxmlElement("w:" + kind + "Reference")
    node.set(qn("w:id"), "1")
    reference._r.append(node)
    paragraph.add_run(" after")

    root = OxmlElement("w:" + kind + "s")
    for note_id, separator in ((-1, "separator"), (0, "continuationSeparator")):
        item = OxmlElement("w:" + kind)
        item.set(qn("w:id"), str(note_id))
        item.set(qn("w:type"), separator)
        p, r = OxmlElement("w:p"), OxmlElement("w:r")
        r.append(OxmlElement("w:" + separator))
        p.append(r)
        item.append(p)
        root.append(item)
    item = OxmlElement("w:" + kind)
    item.set(qn("w:id"), "1")
    p, r = OxmlElement("w:p"), OxmlElement("w:r")
    marker = Run(r, paragraph)
    marker.font.name = "Courier New"
    marker.font.size = Pt(9)
    marker.font.color.rgb = RGBColor.from_string("993311")
    marker.bold = True
    _script(marker.font, marker_script)
    # Properties without a common CharacterStyle equivalent also have to live
    # through the projection; position and spacing are typical real examples.
    for tag, value in (("spacing", "4"), ("position", "1")):
        prop = OxmlElement("w:" + tag)
        prop.set(qn("w:val"), value)
        marker._r.get_or_add_rPr().append(prop)
    marker._r.append(OxmlElement("w:" + kind + "Ref"))
    p.append(r)
    text_run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Note text"
    text_run.append(text)
    p.append(text_run)
    item.append(p)
    root.append(item)
    part = StoryPart(PackURI("/word/" + kind + "s.xml"), CT.WML_FOOTNOTES if kind == "footnote" else CT.WML_ENDNOTES, root, word.part.package)
    word.part.relate_to(part, RT.FOOTNOTES if kind == "footnote" else RT.ENDNOTES)
    return word


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
@pytest.mark.parametrize(("body_script", "marker_script"), [
    (None, None), ("superscript", "baseline"), ("baseline", "superscript"), ("subscript", "superscript"),
])
def test_reference_and_marker_keep_independent_effective_properties(tmp_path: Path, kind, body_script, marker_script):
    word = _source(kind, body_script, marker_script)
    for cycle in range(3):
        initial = tmp_path / f"source-{cycle}.docx"
        word.save(initial)
        source = Document.open(initial)
        assert source.save(tmp_path / "unchanged.docx").read_bytes() == initial.read_bytes()
        original = next(block for block in source.sections[0].blocks if isinstance(block, Note))
        assert original.reference_run_properties is not None
        assert original.marker_run_properties is not None
        native = source.save(tmp_path / "middle.hwpx")
        hwpx = HwpxDocument.open(native)
        hwpx.strict_validate()
        model = Document.from_hwpx_document(hwpx)
        note = next(block for block in model.sections[0].blocks if isinstance(block, Note))
        assert note.reference_run_properties == original.reference_run_properties
        assert note.marker_run_properties == original.marker_run_properties
        assert not any(getattr(block, "field_type", None) == "JAKAL_NOTE_STYLE" for block in note.nested_blocks)
        word = WordFile(model.save(tmp_path / "returned.docx"))
        run = word.element.xpath(".//w:r[w:" + kind + "Reference]")[0]
        expected = None if body_script is None else body_script
        align = run.find(qn("w:rPr")).find(qn("w:vertAlign"))
        assert (align.get(qn("w:val")) if align is not None else None) == expected
        assert run.find(qn("w:rPr")).find(qn("w:rStyle")) is None
        reference = Run(run, word.paragraphs[0])
        assert reference.font.name == "Arial"
        assert reference.font.size == Pt(14)
        assert reference.font.color.rgb == RGBColor.from_string("224466")
        assert reference.italic
        relation = next(rel for rel in word.part.rels.values() if rel.reltype.endswith("/" + kind + "s"))
        root = etree.fromstring(relation.target_part.blob)
        ns = {"w": qn("w:p")[1:].split("}")[0]}
        marker = root.xpath(".//w:r[w:" + kind + "Ref]", namespaces=ns)[0]
        properties = marker.find(qn("w:rPr"))
        align = properties.find(qn("w:vertAlign"))
        assert (align.get(qn("w:val")) if align is not None else None) == marker_script
        assert properties.find(qn("w:rFonts")).get(qn("w:ascii")) == "Courier New"
        assert properties.find(qn("w:sz")).get(qn("w:val")) == "18"
        assert properties.find(qn("w:color")).get(qn("w:val")) == "993311"
        assert properties.find(qn("w:b")) is not None
        assert properties.find(qn("w:spacing")).get(qn("w:val")) == "4"
        assert properties.find(qn("w:position")).get(qn("w:val")) == "1"


def test_note_formats_survive_native_field_parameter_normalization(tmp_path: Path):
    word = _source("footnote", "superscript", "baseline")
    source = tmp_path / "source.docx"
    word.save(source)
    original = Document.open(source)
    native = HwpxDocument.open(original.save(tmp_path / "native.hwpx"))
    fields = [field for field in native.fields() if "ReferenceRPr" in field.parameter_map()]
    assert len(fields) == 1
    parameters = fields[0].element.find("hp:parameters", NS)
    for child in list(parameters):
        if child.get("name") not in {"Prop", "Command"}:
            parameters.remove(child)
    parameters.set("cnt", "2")
    native.sections[0].mark_modified()
    reopened = Document.open(native.save(tmp_path / "normalized.hwpx"))
    before = next(block for block in original.sections[0].blocks if isinstance(block, Note))
    after = next(block for block in reopened.sections[0].blocks if isinstance(block, Note))
    assert after.reference_run_properties == before.reference_run_properties
    assert after.marker_run_properties == before.marker_run_properties


def test_native_notes_without_word_metadata_keep_the_existing_default(tmp_path: Path):
    source = Document()
    source.append_paragraph("Native")
    source.append_note("Note", kind="footNote")
    result = WordFile(source.save(tmp_path / "native.docx"))
    assert result.element.xpath('.//w:r[w:footnoteReference]/w:rPr/w:vertAlign[@w:val="superscript"]')
