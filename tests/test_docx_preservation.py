from __future__ import annotations

from pathlib import Path
import zipfile

from docx import Document
from docx.oxml import OxmlElement
import pytest

from jakal_hwpx import HancomDocument, HwpxDocument, Paragraph, Table

pytestmark = pytest.mark.filterwarnings("ignore:DOCX:UserWarning")


def _fixture(tmp_path):
    document = Document()
    p = document.add_paragraph("Normal ")
    p.add_run("Bold").bold = True
    p.add_run(" tail")
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Unknown content control"
    run.append(text)
    p.append(run)
    content.append(p)
    control.append(content)
    document._element.body.insert(1, control)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Cell text"
    path = tmp_path / "source.docx"
    document.save(path)
    with zipfile.ZipFile(path, "a") as package:
        package.writestr("customXml/opaque.xml", b'<opaque xmlns="urn:test">untouched</opaque>')
        package.comment = b"original package comment"
    return path


def test_unmodified_docx_roundtrip_is_byte_identical(tmp_path: Path) -> None:
    source = _fixture(tmp_path)
    model = HancomDocument.open(source)
    result = model.save(tmp_path / "copy.docx")
    assert result.read_bytes() == source.read_bytes()


def test_text_edits_leave_all_other_package_parts_and_run_properties_intact(tmp_path: Path) -> None:
    source = _fixture(tmp_path)
    model = HancomDocument.open(source)
    blocks = model.sections[0].blocks
    bold = next(block for block in blocks if isinstance(block, Paragraph) and block.text == "Bold")
    bold.text = "Edited bold"
    table = next(block for block in blocks if isinstance(block, Table))
    table.cell_texts[0][0] = "Cell edited\tend"
    model.append_paragraph("Appended\nline")
    result = model.save(tmp_path / "edited.docx")
    reopened = Document(result)
    assert reopened.paragraphs[0].text == "Normal Edited bold tail"
    assert reopened.paragraphs[0].runs[1].bold is True
    assert reopened.tables[0].cell(0, 0).text == "Cell edited\tend"
    assert reopened.paragraphs[-1].text == "Appended\nline"
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(result) as after:
        assert before.namelist() == after.namelist()
        assert before.comment == after.comment
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name), name
        assert b"Unknown content control" in after.read("word/document.xml")


@pytest.mark.parametrize("new", ["", "Normal Bold tail extra", "Completely changed", "A\tB\nC", "Normal bold tail"])
def test_edits_across_styled_run_boundaries_are_not_duplicated(tmp_path: Path, new: str) -> None:
    source = _fixture(tmp_path)
    model = HancomDocument.open(source)
    paragraphs = [b for b in model.sections[0].blocks if isinstance(b, Paragraph)]
    paragraphs[0].text = new
    for paragraph in paragraphs[1:]:
        paragraph.text = ""
    result = model.save(tmp_path / "edited.docx")
    assert Document(result).paragraphs[0].text == new


def test_unsupported_mutation_fails_before_overwriting_output(tmp_path: Path) -> None:
    source = _fixture(tmp_path)
    model = HancomDocument.open(source)
    model.sections[0].settings.page_width += 100
    output = tmp_path / "keep.docx"
    output.write_bytes(b"keep existing output")
    with pytest.raises(ValueError, match="section/header/footer"):
        model.save(output)
    assert output.read_bytes() == b"keep existing output"


def test_preserving_hwpx_converter_receives_the_edited_original_package(tmp_path: Path) -> None:
    source = _fixture(tmp_path)
    model = HancomDocument.open(source)
    model.sections[0].blocks[0].text = "Changed "
    calls = []

    def converter(input_path, output_path, format_name):
        assert format_name == "HWPX"
        assert Document(input_path).paragraphs[0].text == "Changed Bold tail"
        with zipfile.ZipFile(input_path) as package:
            assert package.read("customXml/opaque.xml") == b'<opaque xmlns="urn:test">untouched</opaque>'
        calls.append(str(input_path))
        result = HwpxDocument.blank()
        result.append_paragraph("Native result")
        return result.save(output_path)

    output = model.save(tmp_path / "preserved.hwpx", preserve_formatting=True, converter=converter)
    assert "Native result" in HwpxDocument.open(output).get_document_text()
    assert len(calls) == 1


def test_blank_model_can_export_portable_docx(tmp_path: Path) -> None:
    output = HancomDocument.blank().save(tmp_path / "blank.docx")
    assert len(Document(output).sections) == 1


def test_common_character_style_edits_patch_original_runs(tmp_path: Path) -> None:
    source = _fixture(tmp_path)
    model = HancomDocument.open(source)
    bold = next(style for style in model.character_styles if style.bold)
    bold.italic = True
    bold.height = 1400
    bold.text_color = "#225599"
    bold.font_names["latin"] = "Arial"
    output = model.save(tmp_path / "style-edited.docx")
    run = Document(output).paragraphs[0].runs[1]
    assert run.bold and run.italic
    assert run.font.size.pt == 14
    assert str(run.font.color.rgb) == "225599"
    assert run.font.name == "Arial"
