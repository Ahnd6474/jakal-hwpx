from __future__ import annotations

import base64
from copy import deepcopy
import io
import json
import os
import zipfile
from pathlib import Path
from types import SimpleNamespace
from lxml import etree

from docx import Document as WordFile
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import pytest

from jakal_hwpx import Document, HwpDocument, HwpxDocument, Paragraph, WordDocument
from jakal_hwpx._hancom_models import CharacterStyle, ParagraphStyle
from jakal_hwpx._hancom_styles import _apply_hwp_character_style_payload, _apply_hwp_paragraph_style_payload
from jakal_hwpx.hwp_binary import TAG_PARA_CHAR_SHAPE
from jakal_hwpx.namespaces import NS, qname


PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII=")


@pytest.mark.parametrize("suffix", ["hwpx", "hwp"])
def test_hancom_input_exports_editable_word_content(tmp_path: Path, suffix: str) -> None:
    source = Document.blank()
    source.metadata.title = "Export title"
    source.sections[0].settings.set_page_size(width=61200, height=79200)
    source.sections[0].settings.set_margins(left=7200, right=7200, top=2160, header=5040, bottom=2160, footer=5040)
    source.append_paragraph("Before\t한글\nsecond line")
    source.append_table(rows=2, cols=2, cell_texts=[["Merged", ""], ["Left", "Right"]],
                        col_widths=[14000, 22000], cell_spans={(0, 0): (1, 2)})
    source.append_picture("pixel.png", PNG, extension="png", width=7200, height=3600)
    source.append_hyperlink("https://example.com", display_text="Example link")
    source.append_header("Header")
    source.append_footer("Footer")
    path = source.save(tmp_path / ("source." + suffix))
    document = Document.open(path)
    document.replace_text("Before", "After")
    result = WordFile(document.save(tmp_path / "result.docx"))
    assert "After\t한글\nsecond line" in "\n".join(p.text for p in result.paragraphs)
    assert result.tables[0].cell(0, 0)._tc is result.tables[0].cell(0, 1)._tc
    assert result.tables[0].cell(1, 1).text == "Right"
    assert len(result.inline_shapes) == 1
    assert result.inline_shapes[0].width == Inches(1)
    assert result.inline_shapes[0].height == Inches(0.5)
    assert "Example link" in "\n".join(p.text for p in result.paragraphs)
    assert result.sections[0].header.paragraphs[0].text == "Header"
    assert result.sections[0].footer.paragraphs[0].text == "Footer"
    assert result.sections[0].page_width == Inches(8.5)
    assert result.sections[0].page_height == Inches(11)
    assert result.sections[0].top_margin == Inches(1)
    # The specialized result remains editable through the same shared API.
    word = document.to_docx_document()
    assert isinstance(word, WordDocument)
    word.append_paragraph("Editable")
    assert WordFile(word.save(tmp_path / "editable.docx")).paragraphs[-1].text == "Editable"


def test_hwpx_mixed_runs_and_rich_cells_survive_export(tmp_path: Path) -> None:
    source = WordFile()
    paragraph = source.add_paragraph("Normal ")
    paragraph.add_run("Bold").bold = True
    paragraph.add_run(" italic").italic = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(12)
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("Cell bold").bold = True
    table.cell(0, 1).text = "Other"
    path = tmp_path / "source.docx"
    source.save(path)
    model = Document.open(path)
    hwpx = model.save(tmp_path / "source.hwpx")
    # The DOCX -> common adapter explicitly projects cells as plain text. Set
    # rich content on the HWPX input itself to test this export direction.
    native = HwpxDocument.open(hwpx)
    bold_id = next(style.style_id for style in model.character_styles if style.bold)
    native.tables()[0].cell(0, 0).element.xpath(".//hp:run", namespaces={"hp": "http://www.hancom.co.kr/hwpml/2011/paragraph"})[0].set("charPrIDRef", bold_id)
    native.sections[0].mark_modified()
    native.save(hwpx)
    result = WordFile(Document.open(hwpx).save(tmp_path / "result.docx"))
    paragraph = next(p for p in result.paragraphs if "Normal" in p.text)
    assert paragraph.text == "Normal Bold italic"
    assert paragraph.runs[1].bold is True
    assert paragraph.runs[2].italic is True
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph.paragraph_format.space_after == Pt(12)
    assert result.tables[0].cell(0, 0).paragraphs[0].runs[0].bold is True


@pytest.mark.parametrize("suffix", ["hwpx", "hwp"])
def test_low_level_docx_export_includes_live_edits(tmp_path: Path, suffix: str) -> None:
    source = Document.blank()
    source.append_paragraph("Source")
    path = source.save(tmp_path / ("source." + suffix))
    document = HwpxDocument.open(path) if suffix == "hwpx" else HwpDocument.open(path)
    document.append_paragraph("Live edit")
    output = document.save_as_docx(tmp_path / "live.docx")
    assert "Live edit" in "\n".join(p.text for p in WordFile(output).paragraphs)


def test_docx_export_dispatches_converter_with_current_model_and_preserves_output_on_failure(tmp_path: Path) -> None:
    document = Document.blank()
    document.append_paragraph("Current edit")
    destination = tmp_path / "result.docx"
    calls = []

    def converter(source, target, output_format):
        assert "Current edit" in HwpxDocument.open(source).get_document_text()
        calls.append(output_format)
        result = WordFile()
        result.add_paragraph("External conversion")
        result.save(target)
        return target

    document.save(destination, converter=converter)
    assert calls == ["DOCX"]
    assert WordFile(destination).paragraphs[0].text == "External conversion"
    original = destination.read_bytes()

    def failed(source, target, output_format):
        raise RuntimeError("Conversion failed")

    with pytest.raises(RuntimeError, match="Conversion failed"):
        document.save(destination, converter=failed)
    assert destination.read_bytes() == original
    with pytest.raises(ValueError, match="explicit converter"):
        document.save(destination, preserve_formatting=True)


def test_unsupported_equations_are_not_silently_discarded(tmp_path: Path) -> None:
    document = Document.blank()
    document.append_equation("matrix {x & y}")
    destination = tmp_path / "equation.docx"
    destination.write_bytes(b"previous output")
    with pytest.raises(ValueError, match="equations require an external"):
        document.save(destination)
    assert destination.read_bytes() == b"previous output"


def test_hwpx_cell_text_edit_supersedes_retained_rich_content(tmp_path: Path) -> None:
    source = Document.blank()
    source.append_table(rows=1, cols=1, cell_texts=[["Before"]])
    path = source.save(tmp_path / "source.hwpx")
    document = Document.open(path)
    next(block for block in document.sections[0].blocks if not isinstance(block, Paragraph)).cell_texts[0][0] = "After"
    result = WordFile(document.save(tmp_path / "edited.docx"))
    assert result.tables[0].cell(0, 0).text == "After"


def test_word_to_docx_document_preserves_original_package(tmp_path: Path) -> None:
    source = WordFile()
    source.add_paragraph("Word")
    path = tmp_path / "word.docx"
    source.save(path)
    output = Document.open(path).to_docx_document().save(tmp_path / "copy.docx")
    assert path.read_bytes() == output.read_bytes()


def test_table_fill_border_and_geometry_export(tmp_path: Path) -> None:
    source = Document.blank()
    block = source.append_table(rows=1, cols=2, cell_texts=[["A", "B"]], col_widths=[10000, 20000])
    block.cell_fill_colors = {(0, 0): "#DDEEFF"}
    block.cell_border_styles = {(0, 0): {"left": {"type": "SOLID", "width": "0.3 mm", "color": "#112233"}}}
    hwpx = source.save(tmp_path / "table.hwpx")
    result = WordFile(Document.open(hwpx).save(tmp_path / "table.docx"))
    cell = result.tables[0].cell(0, 0)
    assert cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "DDEEFF"
    assert cell._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:left")).get(qn("w:color")) == "112233"
    assert result.tables[0].columns[1].width == 20000 * 127


def test_hwp_character_styles_and_fonts_export(tmp_path: Path) -> None:
    source = Document.blank()
    source.append_character_style(style_id="9", font_names={"latin": "Arial", "hangul": "맑은 고딕"},
                                  height=1400, bold=True, italic=True)
    source.append_style("Styled", style_id="9", para_pr_id="0", char_pr_id="9")
    source.append_paragraph("Styled 한글").style_id = "9"
    path = source.save(tmp_path / "styled.hwp")
    result = WordFile(Document.open(path).save(tmp_path / "styled.docx"))
    run = next(p for p in result.paragraphs if "Styled" in p.text).runs[0]
    assert run.bold is True
    assert run.italic is True
    assert run.font.name == "Arial"
    assert run.font.size == Pt(14)
    assert run._r.rPr.rFonts.get(qn("w:eastAsia")) == "맑은 고딕"


def test_long_hwp_preview_metadata_does_not_block_word_export(tmp_path: Path) -> None:
    source = Document.blank()
    source.metadata.title = "한글 preview " * 100
    source.append_paragraph("Body")
    result = WordFile(source.save(tmp_path / "long-title.docx"))
    assert result.core_properties.title == source.metadata.title


@pytest.mark.parametrize("suffix", ["hwpx", "hwp"])
def test_external_export_uses_original_until_model_changes(tmp_path: Path, suffix: str) -> None:
    original = Document.blank()
    original.append_paragraph("Original")
    path = original.save(tmp_path / ("original." + suffix))
    raw = path.read_bytes()
    calls = []

    def converter(source, target, output_format):
        calls.append((source.suffix, source.read_bytes(), output_format))
        result = WordFile()
        result.add_paragraph("Converted")
        result.save(target)
        return target

    document = Document.open(path)
    document.save(tmp_path / "original.docx", converter=converter)
    assert calls[0] == ("." + suffix, raw, "DOCX")
    document.replace_text("Original", "Edited")
    document.save(tmp_path / "edited.docx", converter=converter)
    assert calls[1][0] == ".hwpx"
    with zipfile.ZipFile(io.BytesIO(calls[1][1])) as package:
        assert b"Edited" in package.read("Contents/section0.xml")


@pytest.mark.skipif(os.name != "nt", reason="Windows-specific process invocation")
def test_hancom_docx_uses_ooxml_filter(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from jakal_hwpx import hancom_converter

    source = tmp_path / "source.hwpx"
    source.write_bytes(b"source")

    def invoke(arguments, **kwargs):
        config = json.loads(Path(arguments[-1]).read_text(encoding="utf-8"))
        assert config["format"] == "OOXML"
        assert Path(config["output_path"]).suffix == ".docx"
        Path(config["output_path"]).write_bytes(b"converted")
        return SimpleNamespace(returncode=0, stderr=b"")

    monkeypatch.setattr(hancom_converter.subprocess, "run", invoke)
    result = hancom_converter.HancomOfficeConverter()(source, tmp_path / "result.docx", "DOCX")
    assert result.read_bytes() == b"converted"


def test_native_hwp_alignment_distances_and_colorref_use_file_specification(tmp_path: Path) -> None:
    source = Document.blank()
    char_payload = bytearray(74)
    char_payload[14:21] = bytes([100] * 7)
    char_payload[42:46] = (1200).to_bytes(4, "little")
    char_payload[52:56] = bytes.fromhex("22559900")  # native COLORREF: rr gg bb 00
    char_payload[60:64] = bytes.fromhex("ffff0000")  # yellow, not cyan
    character = CharacterStyle(style_id="9")
    _apply_hwp_character_style_payload(character, char_payload)
    source.character_styles.append(character)
    para_payload = bytearray(54)
    para_payload[:4] = (2 << 2).to_bytes(4, "little")  # RIGHT in bits 2..4
    para_payload[4:8] = (2400).to_bytes(4, "little", signed=True)  # native half-units -> 12 pt
    para_payload[12:16] = (600).to_bytes(4, "little", signed=True)
    para_payload[20:24] = (1200).to_bytes(4, "little", signed=True)
    para_payload[46:50] = (1).to_bytes(4, "little")  # fixed line spacing
    para_payload[50:54] = (1400).to_bytes(4, "little")
    paragraph_style = ParagraphStyle(style_id="9")
    _apply_hwp_paragraph_style_payload(paragraph_style, para_payload)
    source.paragraph_styles.append(paragraph_style)
    block = source.append_paragraph("Native formatting")
    block.char_pr_id = block.para_pr_id = "9"
    result = WordFile(source.save(tmp_path / "native.docx"))
    paragraph = result.paragraphs[0]
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paragraph.paragraph_format.left_indent == Pt(12)
    assert paragraph.paragraph_format.first_line_indent == Pt(3)
    assert paragraph.paragraph_format.space_after == Pt(6)
    assert paragraph.paragraph_format.line_spacing == Pt(14)
    assert str(paragraph.runs[0].font.color.rgb) == "225599"
    assert paragraph.runs[0]._r.rPr.find(qn("w:shd")).get(qn("w:fill")) == "FFFF00"
    paragraph_style.alignment_horizontal = "CENTER"
    character.text_color = "#AA1122"
    result = WordFile(source.save(tmp_path / "native-edited.docx"))
    assert result.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert str(result.paragraphs[0].runs[0].font.color.rgb) == "AA1122"


def test_native_hwp_rich_cell_uses_character_runs_and_border_fill(tmp_path: Path) -> None:
    source = Document.blank()
    source.append_character_style(style_id="9", font_names={"latin": "Arial"}, height=1000, bold=True)
    source.append_table(rows=1, cols=1, cell_texts=[["Styled cell"]])
    native = source.to_hwp_document()
    table = native.section(0).tables()[0]
    section_model = native.section_model(0)
    paragraph = next(item.header for item in section_model.paragraphs() if item.header.level > 0)
    shape = next(child for child in paragraph.children if child.tag_id == TAG_PARA_CHAR_SHAPE)
    shape.payload = bytes(4) + (9).to_bytes(4, "little")
    native.binary_document().replace_section_model(0, section_model)
    docinfo = native.docinfo_model()
    record = docinfo.border_fill_records()[table.cells()[0].border_fill_id - 1]
    payload = bytearray(record.payload[:32]) + bytearray(21)
    payload[2] = 1  # solid left border
    payload[4:8] = bytes.fromhex("11223300")
    payload[32:36] = (1).to_bytes(4, "little")  # solid fill
    payload[36:40] = bytes.fromhex("ddeeff00")
    record.payload = bytes(payload)
    native.binary_document().replace_docinfo_model(docinfo)
    path = native.save(tmp_path / "native-cell.hwp")
    result = WordFile(Document.open(path).save(tmp_path / "native-cell.docx"))
    cell = result.tables[0].cell(0, 0)
    assert cell.paragraphs[0].runs[0].bold is True
    assert cell.paragraphs[0].runs[0].font.name == "Arial"
    assert cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) == "DDEEFF"
    assert cell._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:left")).get(qn("w:color")) == "112233"


@pytest.mark.parametrize("modern", [True, False])
def test_hwpx_margin_switch_uses_supported_case_or_legacy_default(tmp_path: Path, modern: bool) -> None:
    source = Document.blank()
    source.append_paragraph_style(style_id="9", margins={"next": "1200"})
    source.append_paragraph("Switch distances").para_pr_id = "9"
    native = source.to_hwpx_document()
    style = native.header.root_element.xpath(".//hh:paraPr[@id='9']", namespaces=NS)[0]
    margin = style.find(qname("hh", "margin"))
    switch = etree.SubElement(style, qname("hp", "switch"))
    case = etree.SubElement(switch, qname("hp", "case"))
    case.set(qname("hp", "required-namespace"), "http://www.hancom.co.kr/hwpml/2016/HwpUnitChar" if modern else "urn:unsupported")
    case.append(deepcopy(margin))
    fallback = etree.SubElement(switch, qname("hp", "default"))
    fallback.append(deepcopy(margin))
    fallback.find("./" + qname("hh", "margin") + "/" + qname("hc", "next")).set("value", "2400")
    style.remove(margin)
    document = Document.from_hwpx_document(native)
    result = WordFile(document.save(tmp_path / "switch.docx"))
    assert result.paragraphs[0].paragraph_format.space_after == Pt(12)
    # Explicit shared edits take precedence over retained native formatting.
    next(item for item in document.paragraph_styles if item.style_id == "9").margins["next"] = "600"
    edited = WordFile(document.save(tmp_path / "edited.docx"))
    assert edited.paragraphs[0].paragraph_format.space_after == Pt(6)
