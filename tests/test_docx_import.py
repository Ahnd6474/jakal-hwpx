from __future__ import annotations

import builtins
import base64
import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from jakal_hwpx import HancomDocument, HwpDocument, HwpHwpxBridge, HwpxDocument, Paragraph, Picture, Table, WordDocument
from jakal_hwpx.namespaces import NS
from jakal_hwpx.hwp_binary import TAG_FACE_NAME


pytestmark = pytest.mark.filterwarnings("ignore:DOCX import:UserWarning")


def _save(document, tmp_path: Path, name: str = "input.docx") -> Path:
    path = tmp_path / name
    document.save(path)
    return path


def _add_hyperlink(paragraph, text: str, target: str = "https://example.com", anchor: str | None = None):
    hyperlink = OxmlElement("w:hyperlink")
    if target:
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    if anchor:
        hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _paragraph_text(section) -> list[str]:
    return [block.text for block in section.blocks if isinstance(block, Paragraph)]


def test_docx_uses_existing_model_and_can_edit_and_export(tmp_path: Path) -> None:
    source = Document()
    source.core_properties.title = "Word 보고서"
    source.core_properties.author = "작성자"
    source.core_properties.keywords = "docx, hancom"
    source.add_paragraph("첫 문단\t탭\n줄바꿈")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(1, 1).text = "B"
    source.add_paragraph("")
    source.add_paragraph("마지막 문단")
    path = _save(source, tmp_path)

    document = HancomDocument.open(path)
    assert isinstance(document, HancomDocument)
    assert type(document) is WordDocument
    assert document.source_format == "docx"
    assert document.metadata.title == "Word 보고서"
    assert document.metadata.creator == "작성자"
    assert document.metadata.keyword == "docx, hancom"
    assert [type(block) for block in document.sections[0].blocks] == [Paragraph, Table, Paragraph, Paragraph]
    assert _paragraph_text(document.sections[0]) == ["첫 문단\t탭\n줄바꿈", "", "마지막 문단"]
    assert document.sections[0].blocks[1].cell_texts == [["A", ""], ["", "B"]]

    document.sections[0].blocks[0].text = "수정한 본문"
    document.append_paragraph("추가한 본문")
    output = document.save(tmp_path / "output.hwpx", preserve_formatting=False)
    reopened = HwpxDocument.open(output)
    assert reopened.metadata().title == "Word 보고서"
    assert "수정한 본문" in reopened.get_document_text()
    assert "추가한 본문" in reopened.get_document_text()
    assert reopened.tables()[0].cell(1, 1).text == "B"


def test_docx_sections_and_inherited_headers(tmp_path: Path) -> None:
    source = Document()
    first = source.sections[0]
    first.page_width, first.page_height = Inches(8), Inches(10)
    first.left_margin = Inches(1)
    first.header.paragraphs[0].text = "공통 머리말"
    first.footer.paragraphs[0].text = "첫 꼬리말"
    source.add_paragraph("첫 섹션")
    second = source.add_section(WD_SECTION_START.NEW_PAGE)
    second.orientation = WD_ORIENT.LANDSCAPE
    second.page_width, second.page_height = Inches(10), Inches(8)
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = "둘째 꼬리말"
    source.add_paragraph("둘째 섹션")

    document = HancomDocument.read_docx(_save(source, tmp_path))
    assert len(document.sections) == 2
    assert "첫 섹션" in _paragraph_text(document.sections[0])
    assert _paragraph_text(document.sections[0]) == ["첫 섹션"]
    assert "둘째 섹션" not in _paragraph_text(document.sections[0])
    assert _paragraph_text(document.sections[1]) == ["둘째 섹션"]
    assert document.sections[0].settings.page_width == 57600
    assert document.sections[0].settings.margins["left"] == 7200
    assert document.sections[1].settings.landscape == "LANDSCAPE"
    assert document.sections[1].settings.page_height == 57600
    assert [block.text for block in document.sections[0].header_footer_blocks] == ["공통 머리말", "첫 꼬리말"]
    assert [block.text for block in document.sections[1].header_footer_blocks] == ["공통 머리말", "둘째 꼬리말"]
    reopened = HwpxDocument.open(document.save(tmp_path / "sections.hwpx", preserve_formatting=False))
    assert len(reopened.sections) == 2
    page = reopened.sections[1].root_element.xpath(".//hp:pagePr", namespaces=NS)[0]
    assert (page.get("landscape"), page.get("width"), page.get("height")) == ("NARROWLY", "57600", "72000")
    assert [header.text for header in reopened.headers()] == ["공통 머리말", "공통 머리말"]


def test_docx_preserves_mixed_run_colors_and_paragraph_style_in_hwpx(tmp_path: Path) -> None:
    source = Document()
    parent = source.styles.add_style("Base", WD_STYLE_TYPE.PARAGRAPH)
    parent.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parent.paragraph_format.space_after = Pt(6)
    parent.font.size = Pt(14)
    derived = source.styles.add_style("Derived", WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = parent
    paragraph = source.add_paragraph(style=derived)
    paragraph.add_run("빨강").font.color.rgb = RGBColor(255, 0, 0)
    paragraph.add_run("파랑").font.color.rgb = RGBColor(0, 0, 255)

    document = HancomDocument.read_docx(_save(source, tmp_path))
    run_ids = {block.char_pr_id for block in document.sections[0].blocks if isinstance(block, Paragraph)}
    run_styles = [style for style in document.character_styles if style.style_id in run_ids]
    assert [style.height for style in run_styles] == [1400, 1400]
    assert [style.text_color for style in run_styles] == ["#FF0000", "#0000FF"]
    assert document.paragraph_styles[0].alignment_horizontal == "CENTER"
    assert document.paragraph_styles[0].margins["next"] == "600"
    assert document.style_definitions[0].name == "Derived"
    reopened = HwpxDocument.open(document.save(tmp_path / "styles.hwpx", preserve_formatting=False))
    runs = reopened.sections[0].root_element.xpath("./hp:p/hp:run[hp:t/text()]", namespaces=NS)
    assert [run.xpath("string(hp:t)", namespaces=NS) for run in runs] == ["빨강", "파랑"]
    assert [run.get("charPrIDRef") for run in runs] == [style.style_id for style in run_styles]
    hwp = HwpDocument.open(document.save(tmp_path / "styles.hwp"))
    hwp_styles = hwp.docinfo_model().char_shape_records()
    assert [hwp_styles[int(style.style_id)].fields()["text_color"] for style in run_styles] == [
        "#FF0000", "#0000FF"
    ]


def test_docx_character_width_is_written_to_both_native_formats(tmp_path: Path) -> None:
    source = Document()
    run = source.add_paragraph().add_run("Narrow")
    run.font.name = "Codex Test Serif"
    width = OxmlElement("w:w")
    width.set(qn("w:val"), "80")
    run._r.get_or_add_rPr().append(width)
    document = HancomDocument.read_docx(_save(source, tmp_path))
    char_id = next(block.char_pr_id for block in document.sections[0].blocks if isinstance(block, Paragraph) and block.text == "Narrow")
    style = next(style for style in document.character_styles if style.style_id == char_id)
    assert style.font_ratios["latin"] == 80

    hwpx = HwpxDocument.open(document.save(tmp_path / "ratio.hwpx"))
    ratios = hwpx.header.root_element.xpath(
        f".//hh:charPr[@id='{style.style_id}']/hh:ratio/@latin", namespaces=NS
    )
    assert ratios == ["80"]
    hwp = HwpDocument.open(document.save(tmp_path / "ratio.hwp"))
    docinfo = hwp.docinfo_model()
    fields = docinfo.char_shape_records()[int(style.style_id)].fields()
    assert fields["ratios"]["latin"] == 80
    latin_id = fields["font_refs"]["latin"]
    latin_start = docinfo.id_mappings_record().counts[1]
    assert docinfo.records_by_tag_id(TAG_FACE_NAME)[latin_start + latin_id].fields()["name"] == "Codex Test Serif"


def test_docx_links_and_picture_keep_inline_order(tmp_path: Path) -> None:
    source = Document()
    paragraph = source.add_paragraph("앞 ")
    _add_hyperlink(paragraph, "링크", anchor="section")
    paragraph.add_run(" 중간 ")
    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII="
    )
    paragraph.add_run().add_picture(io.BytesIO(png), width=Inches(1), height=Inches(0.5))
    paragraph.add_run(" 뒤")
    document = HancomDocument.open(_save(source, tmp_path))
    blocks = document.sections[0].blocks
    assert [type(block).__name__ for block in blocks] == ["Paragraph", "Hyperlink", "Paragraph", "Picture", "Paragraph"]
    assert blocks[1].target == "https://example.com#section"
    picture = next(block for block in blocks if isinstance(block, Picture))
    assert picture.data == png
    assert (picture.width, picture.height) == (7200, 3600)
    reopened = HwpxDocument.open(document.save(tmp_path / "inline.hwpx", preserve_formatting=False))
    assert len(reopened.sections[0].paragraphs()) == 1
    assert reopened.get_document_text().strip() == "앞 링크 중간  뒤"
    assert reopened.pictures()[0].size()["width"] == 7200
    assert reopened.pictures()[0].binary_data() == png
    assert reopened.hyperlinks()[0].hyperlink_target.startswith("https://example.com")


def test_docx_table_merges_and_dimensions_survive_hwpx_export(tmp_path: Path) -> None:
    source = Document()
    table = source.add_table(rows=3, cols=3)
    table.columns[0].width = Inches(1)
    table.columns[1].width = Inches(2)
    table.columns[2].width = Inches(1)
    for row in table.rows:
        row.height = Inches(0.5)
    table.cell(0, 0).merge(table.cell(1, 1)).text = "병합"
    table.cell(2, 2).text = "끝"
    table.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM

    document = HancomDocument.read_docx(_save(source, tmp_path))
    block = document.sections[0].blocks[0]
    assert isinstance(block, Table)
    assert block.cell_spans == {(0, 0): (2, 2)}
    assert block.cell_texts == [["병합", "", ""], ["", "", ""], ["", "", "끝"]]
    assert block.col_widths == [7200, 14400, 7200]
    reopened = HwpxDocument.open(document.save(tmp_path / "merged.hwpx", preserve_formatting=False))
    cell = reopened.tables()[0].cell(0, 0)
    assert (cell.row_span, cell.col_span) == (2, 2)
    assert cell.vertical_align == "BOTTOM"
    assert cell.element.xpath("./hp:cellSz/@width", namespaces=NS) == ["21600"]
    assert cell.element.xpath("./hp:cellSz/@height", namespaces=NS) == ["7200"]
    assert reopened.tables()[0].cell(2, 2).text == "끝"


def test_docx_table_cell_padding_survives_python_exports(tmp_path: Path) -> None:
    source = Document()
    table = source.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Padded"
    table_margins = OxmlElement("w:tblCellMar")
    left = OxmlElement("w:left")
    left.set(qn("w:w"), "120")
    left.set(qn("w:type"), "dxa")
    table_margins.append(left)
    table._tbl.tblPr.append(table_margins)
    cell_margins = OxmlElement("w:tcMar")
    top = OxmlElement("w:top")
    top.set(qn("w:w"), "80")
    top.set(qn("w:type"), "dxa")
    cell_margins.append(top)
    table.cell(0, 0)._tc.get_or_add_tcPr().append(cell_margins)

    document = HancomDocument.read_docx(_save(source, tmp_path))
    block = document.sections[0].blocks[0]
    assert isinstance(block, Table)
    assert block.table_margins == {"left": 600, "right": 540, "top": 0, "bottom": 0}
    assert block.cell_margins == {(0, 0): {"left": 600, "right": 540, "top": 400, "bottom": 0}}
    hwpx = HwpxDocument.open(document.save(tmp_path / "padding.hwpx"))
    assert hwpx.tables()[0].cell(0, 0).margins["top"] == 400
    assert hwpx.tables()[0].cell(0, 0).element.get("hasMargin") == "1"
    hwp = HwpDocument.open(document.save(tmp_path / "padding.hwp"))
    assert hwp.tables()[0].cells()[0].margins[0] == 600
    assert hwp.tables()[0].cells()[0].margins[2] == 400


def test_docx_warns_about_features_outside_the_common_model(tmp_path: Path) -> None:
    source = Document()
    paragraph = source.add_paragraph("목록", style="List Bullet")
    paragraph.add_run("굵게").bold = True
    source.sections[0].different_first_page_header_footer = True
    source.sections[0].first_page_header.paragraphs[0].text = "첫 페이지"
    path = _save(source, tmp_path)
    document = HancomDocument.read_docx(path)
    with pytest.warns(UserWarning) as captured:
        document.to_common_document()
    messages = [str(warning.message) for warning in captured]
    assert messages == document.import_warnings
    assert len(messages) == len(set(messages))
    assert any("automatic numbering" in message for message in messages)
    assert any(style.bold for style in document.character_styles)
    assert any(block.apply_page_type == "FIRST" for block in document.sections[0].header_footer_blocks)


def test_docx_optional_dependency_has_install_instructions(tmp_path: Path, monkeypatch) -> None:
    original_import = builtins.__import__

    def without_docx(name, *args, **kwargs):
        if name == "docx":
            raise ModuleNotFoundError("No module named 'docx'", name="docx")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_docx)
    with pytest.raises(ModuleNotFoundError, match=r"jakal-hwpx\[docx\]"):
        HancomDocument.read_docx(tmp_path / "input.docx")
    assert HancomDocument.blank().source_format is None


def test_docx_sniffing_and_bridge_and_unsupported_output(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("공통 모델")
    path = _save(source, tmp_path, "document.bin")
    assert HancomDocument.open(path).source_format == "docx"
    def converter(source_path, target_path, format_name):
        doc = HwpxDocument.blank()
        doc.append_paragraph(Document(source_path).paragraphs[0].text)
        return doc.save(target_path)

    bridge = HwpHwpxBridge.open(path, converter=converter)
    assert bridge.hancom_document().source_format == "docx"
    assert "공통 모델" in bridge.hwpx_document().get_document_text()
    output = HancomDocument.open(path).save(tmp_path / "output.docx")
    assert output.read_bytes() == path.read_bytes()
    path.write_bytes(b"not a docx")
    from docx.opc.exceptions import PackageNotFoundError

    with pytest.raises(PackageNotFoundError):
        HancomDocument.read_docx(path)


def test_docx_subclass_factory_and_empty_document(tmp_path: Path) -> None:
    class CustomDocument(HancomDocument):
        pass

    path = _save(Document(), tmp_path)
    document = CustomDocument.open(path)
    assert isinstance(document, CustomDocument)
    assert len(document.sections) == 1
    assert document.sections[0].blocks == []
    assert document.to_hwpx_document(preserve_formatting=False).strict_validate() is None


def test_docx_document_defaults_are_inherited_without_a_named_style_override(tmp_path: Path) -> None:
    source = Document()
    defaults = source.styles.element.xpath("./w:docDefaults/w:pPrDefault/w:pPr")[0]
    spacing = defaults.find(qn("w:spacing"))
    spacing.set(qn("w:after"), "320")
    spacing.set(qn("w:line"), "360")
    spacing.set(qn("w:lineRule"), "auto")
    source.add_paragraph("Defaults")
    document = HancomDocument.read_docx(_save(source, tmp_path))
    block = document.sections[0].blocks[0]
    style = next(style for style in document.paragraph_styles if style.style_id == block.para_pr_id)
    assert style.margins["next"] == "1600"
    assert (style.line_spacing, style.line_spacing_type) == (150, "PERCENT")
    reopened = document.to_hwpx_document(preserve_formatting=False)
    style = next(style for style in reopened.paragraph_styles() if style.style_id == block.para_pr_id)
    assert style.element.xpath(".//hh:lineSpacing/@value", namespaces=NS) == ["150"]


def test_docx_headers_do_not_add_body_paragraphs(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("Body")
    source.sections[0].header.paragraphs[0].text = "Header"
    source.sections[0].footer.paragraphs[0].text = "Footer"
    document = HancomDocument.read_docx(_save(source, tmp_path))
    reopened = document.to_hwpx_document(preserve_formatting=False)
    assert len(reopened.sections[0].paragraphs()) == 1
    assert reopened.headers()[0].text == "Header"
    assert reopened.footers()[0].text == "Footer"
    assert reopened.sections[0].root_element.xpath("./hp:p[1]/hp:run[hp:secPr]/hp:ctrl/hp:colPr/@colCount", namespaces=NS) == ["1"]


def test_docx_margins_keep_body_offsets_and_header_footer_positions(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("Body")
    section = source.sections[0]
    section.top_margin, section.header_distance = Inches(1.25), Inches(0.25)
    section.bottom_margin, section.footer_distance = Inches(2), Inches(0.5)
    document = HancomDocument.read_docx(_save(source, tmp_path))
    margins = document.sections[0].settings.margins
    assert (margins["top"], margins["header"]) == (1800, 7200)
    assert (margins["bottom"], margins["footer"]) == (3600, 10800)
    exported = document.to_hwpx_document(preserve_formatting=False).section_settings().margins()
    assert exported["top"] + exported["header"] == 9000
    assert exported["bottom"] + exported["footer"] == 14400


def test_docx_named_style_does_not_include_first_paragraph_direct_overrides(tmp_path: Path) -> None:
    source = Document()
    source.styles["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    source.add_paragraph("Override").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    source.add_paragraph("Base")
    document = HancomDocument.read_docx(_save(source, tmp_path))
    named = document.style_definitions[0]
    base = next(style for style in document.paragraph_styles if style.style_id == named.para_pr_id)
    override = next(style for style in document.paragraph_styles if style.style_id == document.sections[0].blocks[0].para_pr_id)
    assert base.alignment_horizontal == "LEFT"
    assert override.alignment_horizontal == "RIGHT"


def test_docx_first_page_header_behavior_keeps_an_empty_first_header(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("Body")
    section = source.sections[0]
    section.header.paragraphs[0].text = "Header after first page"
    section.different_first_page_header_footer = True
    document = HancomDocument.read_docx(_save(source, tmp_path))
    assert document.sections[0].settings.visibility["hideFirstHeader"] == "1"
    assert not any("first-page" in message for message in document.import_warnings)


@pytest.mark.parametrize("merges", [
    [(0, 0, 2, 2)],
    [(0, 0, 0, 3), (1, 0, 3, 0), (1, 1, 2, 3)],
    [(0, 0, 3, 3)],
    [(0, 0, 1, 1), (0, 2, 3, 3), (2, 0, 3, 1)],
])
def test_docx_table_merge_matrix_has_no_lost_or_duplicated_text(tmp_path: Path, merges) -> None:
    source = Document()
    table = source.add_table(rows=4, cols=4)
    expected_texts = [[f"r{row}c{col}" for col in range(4)] for row in range(4)]
    expected_spans = {}
    for top, left, bottom, right in merges:
        table.cell(top, left).merge(table.cell(bottom, right))
        expected_spans[(top, left)] = (bottom - top + 1, right - left + 1)
        for row in range(top, bottom + 1):
            for col in range(left, right + 1):
                expected_texts[row][col] = ""
        expected_texts[top][left] = f"merged{top}{left}{bottom}{right}"
    for row in range(4):
        for col in range(4):
            if expected_texts[row][col]:
                table.cell(row, col).text = expected_texts[row][col]
    document = HancomDocument.read_docx(_save(source, tmp_path))
    block = document.sections[0].blocks[0]
    assert block.cell_texts == expected_texts
    assert block.cell_spans == expected_spans
    reopened = HancomDocument.from_hwpx_document(document.to_hwpx_document(preserve_formatting=False))
    block = next(block for block in reopened.sections[0].blocks if isinstance(block, Table))
    assert block.cell_texts == expected_texts
    assert block.cell_spans == expected_spans


@pytest.mark.parametrize("text", ["A\tB\nC", "\tA\t\nB\n", "\n\n\t\t", "한글\t🙂\n끝"])
def test_docx_tab_linebreak_encoding_survives_read_and_edit(tmp_path: Path, text: str) -> None:
    source = Document()
    source.add_paragraph(text)
    source.add_table(rows=1, cols=1).cell(0, 0).text = text
    document = HancomDocument.read_docx(_save(source, tmp_path))
    output = document.save(tmp_path / "controls.hwpx", preserve_formatting=False)
    reopened = HwpxDocument.open(output)
    root = reopened.sections[0].root_element
    body = root.xpath("./hp:p[1]/hp:run/hp:t", namespaces=NS)
    assert sum(len(node.findall("hp:tab", NS)) for node in body) == text.count("\t")
    assert sum(len(node.findall("hp:lineBreak", NS)) for node in body) == text.count("\n")
    for node in root.xpath(".//hp:t", namespaces=NS):
        assert not any(char in (node.text or "") for char in "\t\r\n")
        assert all(not any(char in (child.tail or "") for char in "\t\r\n") for child in node)
    model = HancomDocument.open(output)
    assert model.sections[0].blocks[0].text == text
    table = next(block for block in model.sections[0].blocks if isinstance(block, Table))
    assert table.cell_texts[0][0] == text
    assert text in reopened.get_document_text()
    # Replacement must also find text in the XML tails after a tab/line break.
    if "B" in text:
        assert reopened.replace_text("B", "EDIT\tNEW") == 2
        edited = HancomDocument.from_hwpx_document(reopened)
        assert edited.sections[0].blocks[0].text == text.replace("B", "EDIT\tNEW")
        assert reopened.tables()[0].cell(0, 0).text == text.replace("B", "EDIT\tNEW")


def test_docx_hwpx_replacement_keeps_tab_properties_and_counts_original_matches(tmp_path: Path) -> None:
    source = Document()
    source.add_paragraph("A\tA\nA")
    document = HancomDocument.read_docx(_save(source, tmp_path)).to_hwpx_document(preserve_formatting=False)
    tab = document.sections[0].root_element.xpath(".//hp:t/hp:tab", namespaces=NS)[0]
    tab.set("width", "567")
    tab.set("type", "1")
    document.sections[0].mark_modified()
    assert document.replace_text("A", "AA", count=2) == 2
    assert document.get_document_text() == "AA\tAA\nA"
    tab = document.sections[0].root_element.xpath(".//hp:t/hp:tab", namespaces=NS)[0]
    assert (tab.get("width"), tab.get("type")) == ("567", "1")


def test_docx_landscape_keeps_physical_page_size_across_hwp_hwpx(tmp_path: Path) -> None:
    source = Document()
    section = source.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(10), Inches(8)
    source.add_paragraph("Landscape")
    document = HancomDocument.read_docx(_save(source, tmp_path))
    hwpx = document.to_hwpx_document(preserve_formatting=False)
    through_hwpx = HancomDocument.from_hwpx_document(hwpx).to_hwp_document(preserve_formatting=False)
    direct = document.to_hwp_document(preserve_formatting=False)
    for hwp in (direct, through_hwpx):
        settings = hwp.binary_document().section_page_settings(0)
        assert (settings["page_width"], settings["page_height"]) == (72000, 57600)
