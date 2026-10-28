"""Equivalent display fallbacks must survive a real format boundary."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import base64

from docx import Document as WordFile
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
import pytest

from jakal_hwpx import Document, HwpxDocument
from jakal_hwpx.hancom_document import Field, Note, Paragraph, Picture, Table
from jakal_hwpx.namespaces import NS

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII=")


def _cycle(word, directory):
    initial = directory / "source.docx"
    word.save(initial)
    source = Document.open(initial)
    assert source.save(directory / "unchanged.docx").read_bytes() == initial.read_bytes()
    native_path = source.save(directory / "middle.hwpx")
    native = HwpxDocument.open(native_path)
    native.strict_validate()
    returned_path = Document.open(native_path).save(directory / "returned.docx")
    return WordFile(returned_path), native


def _text(parent, text, *, bold=False):
    run = OxmlElement("w:r")
    if bold:
        properties = OxmlElement("w:rPr")
        properties.append(OxmlElement("w:b"))
        run.append(properties)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    parent.append(run)
    return run


def test_first_page_stories_are_rich_non_wrapping_shapes_and_restore(tmp_path: Path) -> None:
    word = WordFile()
    word.add_paragraph("Body one")
    word.add_page_break()
    word.add_paragraph("Body two")
    section = word.sections[0]
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = "Normal header"
    section.footer.paragraphs[0].text = "Normal footer"
    section.first_page_header.paragraphs[0].add_run("First header").bold = True
    section.first_page_header.paragraphs[0].add_run().add_picture(BytesIO(PNG), width=Inches(.2))
    section.first_page_footer.paragraphs[0].add_run("First footer").italic = True
    for _ in range(3):
        word, native = _cycle(word, tmp_path)
        settings = native.section_settings(0).visibility()
        assert settings["hideFirstHeader"] == settings["hideFirstFooter"] == "1"
        shapes = native.sections[0].readonly_root_element.xpath("./hp:p/hp:run/hp:rect", namespaces=NS)
        assert len(shapes) == 2
        for shape in shapes:
            assert shape.find("hp:pos", NS).get("treatAsChar") == "0"
            assert shape.get("textWrap") == "IN_FRONT_OF_TEXT"
            assert not shape.xpath("./hp:lineShape | ./hc:fillBrush", namespaces=NS)
        section = word.sections[0]
        assert section.different_first_page_header_footer
        assert section.header.paragraphs[0].text == "Normal header"
        assert section.footer.paragraphs[0].text == "Normal footer"
        assert section.first_page_header.paragraphs[0].text == "First header"
        assert section.first_page_header.paragraphs[0].runs[0].bold
        assert section.first_page_header._element.xpath(".//w:drawing")
        assert section.first_page_footer.paragraphs[0].text == "First footer"
        assert section.first_page_footer.paragraphs[0].runs[0].italic
        assert "First header" not in "".join(p.text for p in word.paragraphs)


def test_empty_first_stories_remain_empty(tmp_path: Path) -> None:
    word = WordFile()
    word.add_paragraph("Body")
    word.sections[0].different_first_page_header_footer = True
    word.sections[0].header.paragraphs[0].text = "Later pages"
    word, native = _cycle(word, tmp_path)
    assert word.sections[0].different_first_page_header_footer
    assert not word.sections[0].first_page_header.paragraphs[0].text
    assert not native.shapes()


def test_landscape_first_story_uses_the_oriented_paper_width(tmp_path: Path) -> None:
    from docx.enum.section import WD_ORIENT

    word = WordFile()
    word.add_paragraph("Landscape")
    section = word.sections[0]
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "Wide header"
    _, native = _cycle(word, tmp_path)
    shape = native.shapes()[0]
    expected = round((int(section.page_width) - int(section.left_margin) - int(section.right_margin)) / 127)
    assert shape.size()["width"] == expected


@pytest.mark.parametrize(("kind", "start", "labels"), [
    ("upperRoman", 4, ["IV)", "V)"]),
    ("lowerLetter", 26, ["z)", "aa)"]),
    ("decimalZero", 3, ["03)", "04)"]),
])
def test_list_display_fields_restore_active_word_lists(tmp_path: Path, kind, start, labels) -> None:
    word = WordFile()
    numbering = word.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "20")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", kind), ("lvlText", "%1)"), ("suff", "space")):
        child = OxmlElement("w:" + tag)
        child.set(qn("w:val"), value)
        level.append(child)
    abstract.append(level)
    numbering.insert(0, abstract)
    number = numbering.add_num(20)
    override = number.add_lvlOverride(0)
    override.add_startOverride(start)
    for text in ("first", "second"):
        paragraph = word.add_paragraph(text)
        properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_numId().val = number.numId
        properties.get_or_add_ilvl().val = 0
    for _ in range(3):
        word, native = _cycle(word, tmp_path)
        model = Document.from_hwpx_document(native)
        fields = [block for block in model.sections[0].blocks if isinstance(block, Field)]
        assert [field.display_text for field in fields] == [label + " " for label in labels]
        assert [p.text for p in word.paragraphs] == ["first", "second"]
        ids = [p._p.pPr.numPr.numId.val for p in word.paragraphs]
        assert ids[0] == ids[1]
        target = word.part.numbering_part.element
        number = next(node for node in target if node.tag == qn("w:num") and node.get(qn("w:numId")) == str(ids[0]))
        assert number.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val")) == str(start)


def test_lists_inside_cells_are_not_lost_or_duplicated(tmp_path: Path) -> None:
    word = WordFile()
    table = word.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.style = "List Number"
    paragraph.text = "Nested list"
    word, native = _cycle(word, tmp_path)
    assert "1.\tNested list" in native.get_document_text()
    paragraph = word.tables[0].cell(0, 0).paragraphs[0]
    assert paragraph.text == "Nested list"
    assert paragraph._p.pPr.numPr is not None


def test_native_normalized_fields_restore_after_parameter_loss_and_run_merging(tmp_path: Path) -> None:
    word = WordFile()
    word.add_paragraph("List body", style="List Number")
    _, native = _cycle(word, tmp_path)
    field = native.fields()[0]
    assert field.name.startswith("jakal:word-effect:v1:")
    parameters = field.element.find("hp:parameters", NS)
    for child in list(parameters):
        if child.get("name") not in {"Prop", "Command"}:
            parameters.remove(child)
    parameters.set("cnt", "2")
    paragraph = native.sections[0].root_element.find("hp:p", NS)
    runs = paragraph.findall("hp:run", NS)
    for run in runs[1:]:
        runs[0].extend(list(run))
        paragraph.remove(run)
    native.sections[0].mark_modified()
    assert field.display_text == "1.\t"
    path = native.save(tmp_path / "normalized.hwpx")
    result = WordFile(Document.open(path).save(tmp_path / "normalized.docx"))
    assert result.paragraphs[0].text == "List body"
    assert result.paragraphs[0]._p.pPr.numPr is not None
    # Editing the static label is respected; it must not overwrite nearby body
    # text or silently regenerate a different label from the original definition.
    field.set_display_text("Custom\t")
    assert field.display_text == "Custom\t"
    assert "List body" in native.get_document_text()
    path = native.save(tmp_path / "edited-label.hwpx")
    with pytest.warns(UserWarning, match="Unknown field"):
        result = WordFile(Document.open(path).save(tmp_path / "edited-label.docx"))
    assert result.paragraphs[0].text == "Custom\tList body"
    assert result.paragraphs[0]._p.pPr.numPr is None


@pytest.mark.parametrize("table", [False, True])
def test_standalone_page_break_does_not_create_an_extra_blank_line(tmp_path: Path, table: bool) -> None:
    word = WordFile()
    word.add_paragraph("First page")
    word.add_page_break()
    if table:
        word.add_table(rows=1, cols=1).cell(0, 0).text = "Second page table"
    else:
        word.add_paragraph("Second page")
    result, native = _cycle(word, tmp_path)
    assert len(native.sections[0].readonly_root_element.findall("hp:p", NS)) == 2
    if table:
        assert result.paragraphs[1]._p.xpath('./w:r/w:br[@w:type="page"]')
        assert result.tables[0].cell(0, 0).text == "Second page table"
    else:
        assert [p.text for p in result.paragraphs] == ["First page", "Second page"]
        assert result.paragraphs[1].paragraph_format.page_break_before


def test_fields_bookmarks_and_page_breaks_keep_their_effect(tmp_path: Path) -> None:
    word = WordFile()
    paragraph = word.add_paragraph("Before ")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "1")
    bookmark.set(qn("w:name"), "target")
    paragraph._p.append(bookmark)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'DATE \\@ "yyyy-MM-dd"')
    _text(field, "2026-10-28", bold=True)
    paragraph._p.append(field)
    paragraph.add_run(" after")
    page = word.sections[0].footer.paragraphs[0]
    for tag, value in (("fldChar", "begin"), ("instrText", " PAGE "), ("fldChar", "separate"), ("t", "1"), ("fldChar", "end")):
        run = OxmlElement("w:r")
        item = OxmlElement("w:" + tag)
        if tag == "fldChar":
            item.set(qn("w:fldCharType"), value)
        else:
            item.text = value
        run.append(item)
        page._p.append(run)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("Next page")
    word, native = _cycle(word, tmp_path)
    assert len(native.bookmarks()) == 1
    assert native.sections[0].readonly_root_element.xpath("./hp:p[@pageBreak='1']", namespaces=NS)
    assert native.sections[0].readonly_root_element.xpath(".//hp:autoNum[@numType='PAGE']", namespaces=NS)
    field = word.paragraphs[0]._p.find(qn("w:fldSimple"))
    assert field.get(qn("w:instr")) == 'DATE \\@ "yyyy-MM-dd"'
    assert field.find(qn("w:r")).find(qn("w:rPr")).find(qn("w:b")) is not None
    assert "2026-10-28" not in word.paragraphs[0].text  # native field, not duplicated body text
    assert word.sections[0].footer._element.xpath('.//w:fldSimple[@w:instr="PAGE"]')
    assert word.paragraphs[1].paragraph_format.page_break_before


def test_accepted_revisions_and_content_controls_keep_visible_content_and_source(tmp_path: Path) -> None:
    word = WordFile()
    paragraph = word.add_paragraph("Before ")
    inserted = OxmlElement("w:ins")
    _text(inserted, "accepted", bold=True)
    paragraph._p.append(inserted)
    deleted = OxmlElement("w:del")
    _text(deleted, "deleted")
    paragraph._p.append(deleted)
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    _text(content, " visible")
    control.append(content)
    paragraph._p.append(control)
    table = word.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Wrapped table"
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    table._tbl.addprevious(control)
    control.append(content)
    content.append(table._tbl)
    path = tmp_path / "original.docx"
    word.save(path)
    editable = Document.open(path)
    assert editable.replace_text("accepted", "edited") == 1
    edited = editable.save(tmp_path / "edited.docx")
    assert WordFile(edited).element.xpath(".//w:ins/w:r/w:t[text()='edited']")
    word, native = _cycle(word, tmp_path)
    assert "Before accepted visible" in native.get_document_text()
    assert "deleted" not in native.get_document_text()
    assert word.paragraphs[0].text == "Before accepted visible"
    assert word.paragraphs[0].runs[1].bold
    assert word.tables[0].cell(0, 0).text == "Wrapped table"


def test_hyperlink_inside_content_control_can_still_save_and_edit_original_word(tmp_path: Path) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from jakal_hwpx.hancom_document import Hyperlink

    word = WordFile()
    paragraph = word.add_paragraph("Before ")
    control, content, link = OxmlElement("w:sdt"), OxmlElement("w:sdtContent"), OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to("https://example.com/old", RT.HYPERLINK, is_external=True))
    _text(link, "Link")
    content.append(link)
    control.append(content)
    paragraph._p.append(control)
    path = tmp_path / "wrapped-link.docx"
    word.save(path)
    editable = Document.open(path)
    assert editable.save(tmp_path / "copy.docx").read_bytes() == path.read_bytes()
    block = next(block for block in editable.sections[0].blocks if isinstance(block, Hyperlink))
    block.target = "https://example.com/new"
    result = WordFile(editable.save(tmp_path / "edited.docx"))
    assert result.element.xpath(".//w:sdt/w:sdtContent/w:hyperlink")
    assert any(rel.target_ref == "https://example.com/new" for rel in result.part.rels.values())


@pytest.mark.parametrize("kind", ["footNote", "endNote"])
def test_notes_roundtrip_as_native_references_with_rich_text(tmp_path: Path, kind: str) -> None:
    source = Document()
    source.append_paragraph("Reference")
    note = source.append_note("A note", kind=kind)
    note.nested_blocks = [Paragraph("A "), Paragraph("note")]
    for block in note.nested_blocks:
        setattr(block, "nested_paragraph_index", 0)
    first = source.save(tmp_path / "first.hwpx")
    word_path = Document.open(first).save(tmp_path / "notes.docx")
    word = WordFile(word_path)
    name = "footnote" if kind == "footNote" else "endnote"
    assert word.element.xpath(".//w:" + name + "Reference")
    relationship = next(rel for rel in word.part.rels.values() if rel.reltype.endswith("/" + name + "s"))
    root = etree.fromstring(relationship.target_part.blob)
    assert "A note" in "".join(root.itertext())
    returned = Document.open(word_path).save(tmp_path / "notes.hwpx")
    native = HwpxDocument.open(returned)
    native.strict_validate()
    model = Document.from_hwpx_document(native)
    notes = [block for block in model.sections[0].blocks if isinstance(block, Note)]
    assert len(notes) == 1
    assert notes[0].text == "A note"


def test_note_images_tables_and_character_styles_resolve_in_their_own_part(tmp_path: Path) -> None:
    source = Document()
    source.append_character_style(style_id="1", bold=True)
    source.append_paragraph("Reference")
    note = source.append_note("Rich note", kind="footNote")
    text = Paragraph("Rich note", char_pr_id="1")
    picture = Picture("tiny.png", PNG, "png", width=1440, height=1440)
    table = Table(1, 1, [["Note table"]])
    for block in (text, picture):
        setattr(block, "nested_paragraph_index", 0)
    setattr(table, "nested_paragraph_index", 1)
    note.nested_blocks = [text, picture, table]
    first = source.save(tmp_path / "rich-note.hwpx")
    path = Document.open(first).save(tmp_path / "rich-note.docx")
    word = WordFile(path)
    relationship = next(rel for rel in word.part.rels.values() if rel.reltype.endswith("/footnotes"))
    part = relationship.target_part
    assert any(rel.reltype.endswith("/image") for rel in part.rels.values())
    root = etree.fromstring(part.blob)
    assert root.xpath(".//w:tbl", namespaces={"w": qn("w:p")[1:].split("}")[0]})
    returned = Document.open(path).save(tmp_path / "returned-note.hwpx")
    model = Document.open(returned)
    note = next(block for block in model.sections[0].blocks if isinstance(block, Note))
    assert any(isinstance(block, Picture) and block.data == PNG for block in note.nested_blocks)
    assert any(isinstance(block, Table) and block.cell_texts == [["Note table"]] for block in note.nested_blocks)
    character_ids = {block.char_pr_id for block in note.nested_blocks if isinstance(block, Paragraph) and "Rich note" in block.text}
    assert any(style.style_id in character_ids and style.bold for style in model.character_styles)
