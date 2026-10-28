"""Cross-format contracts, checked with python-docx and native HWPX XML."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import base64

from docx import Document as WordFile
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from lxml import etree
import pytest

from jakal_hwpx import Document, HwpxDocument
from jakal_hwpx.namespaces import NS, qname

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/l9sAAAAASUVORK5CYII=")


def _link(paragraph, text, target="https://example.com"):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    word = OxmlElement("w:t")
    word.text = text
    run.append(word)
    node.append(run)
    paragraph._p.append(node)


def _word_cycle(source, directory):
    initial = directory / "source.docx"
    source.save(initial)
    hwpx = Document.open(initial).save(directory / "middle.hwpx")
    HwpxDocument.open(hwpx).strict_validate()
    result = Document.open(hwpx).save(directory / "result.docx")
    return WordFile(result), HwpxDocument.open(hwpx)


def test_word_hwpx_word_preserves_rich_cells_and_nested_objects(tmp_path: Path) -> None:
    word = WordFile()
    table = word.add_table(rows=2, cols=2)
    paragraph = table.cell(0, 0).paragraphs[0]
    paragraph.add_run("Normal ")
    run = paragraph.add_run("bold red")
    run.bold = True
    run.font.color.rgb = RGBColor(0x22, 0x55, 0x99)
    run.font.size = Pt(14)
    paragraph.add_run("\tline\nbreak")
    _link(paragraph, "Link")
    table.cell(0, 1).paragraphs[0].add_run().add_picture(BytesIO(PNG), width=Inches(0.4), height=Inches(0.2))
    nested = table.cell(1, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested"
    table.cell(1, 1).text = "Last"
    result, native = _word_cycle(word, tmp_path)
    assert native.tables()[0].cell(0, 0).element.xpath(".//hp:run[@charPrIDRef != '0']", namespaces=NS)
    paragraph = result.tables[0].cell(0, 0).paragraphs[0]
    assert paragraph.text == "Normal bold red\tline\nbreakLink"
    run = next(run for run in paragraph.runs if run.text == "bold red")
    assert run.bold is True
    assert run.font.color.rgb == RGBColor(0x22, 0x55, 0x99)
    assert run.font.size == Pt(14)
    assert paragraph.hyperlinks[0].url == "https://example.com"
    assert result.tables[0].cell(1, 0).tables[0].cell(0, 0).text == "Nested"
    image = result.tables[0].cell(0, 1)._tc.xpath(".//wp:extent")[0]
    assert int(image.get("cx")) == Inches(0.4)
    assert int(image.get("cy")) == Inches(0.2)


def test_rich_headers_survive_word_hwpx_word(tmp_path: Path) -> None:
    word = WordFile()
    word.add_paragraph("Body")
    paragraph = word.sections[0].header.paragraphs[0]
    paragraph.add_run("Header ")
    paragraph.add_run("bold").bold = True
    paragraph.add_run().add_picture(BytesIO(PNG), width=Inches(0.2))
    footer = word.sections[0].footer.paragraphs[0]
    footer.add_run("Footer").italic = True
    result, _ = _word_cycle(word, tmp_path)
    header = result.sections[0].header.paragraphs[0]
    assert header.text == "Header bold"
    assert next(run for run in header.runs if run.text == "bold").bold is True
    assert header._p.xpath(".//wp:extent")
    assert result.sections[0].footer.paragraphs[0].runs[0].italic is True


@pytest.mark.parametrize("direction", ["word", "hwpx"])
def test_native_resources_and_paragraph_distances_do_not_drift(tmp_path: Path, direction: str) -> None:
    word = WordFile()
    paragraph = word.add_paragraph("Distances")
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(12)
    paragraph.paragraph_format.first_line_indent = Inches(0.25)
    paragraph.paragraph_format.line_spacing = 1.5
    initial = tmp_path / "source.docx"
    word.save(initial)
    model = Document.open(initial)
    if direction == "hwpx":
        first = model.save(tmp_path / "source.hwpx")
        model = Document.open(first)
    for index in range(3):
        output = model.save(tmp_path / f"cycle-{index}.docx")
        result = WordFile(output)
        paragraph = next(p for p in result.paragraphs if p.text == "Distances")
        assert paragraph.paragraph_format.space_before == Pt(4)
        assert paragraph.paragraph_format.space_after == Pt(12)
        assert paragraph.paragraph_format.first_line_indent == Inches(0.25)
        assert paragraph.paragraph_format.line_spacing == 1.5
        native = Document.open(output).save(tmp_path / f"cycle-{index}.hwpx")
        HwpxDocument.open(native).strict_validate()
        model = Document.open(native)


def test_hwpx_word_hwpx_preserves_native_border_resources_and_switch_distances(tmp_path: Path) -> None:
    common = Document.blank()
    common.append_paragraph("Native body")
    block = common.append_table(rows=1, cols=1, cell_texts=[["Cell"]])
    block.cell_fill_colors = {(0, 0): "#DDEEFF"}
    native = common.to_hwpx_document()
    table = native.tables()[0]
    definition = native.header.root_element.xpath(".//hh:borderFill[@id=$id]", namespaces=NS, id=str(table.cell(0, 0).border_fill_id))[0]
    definition.find(qname("hh", "leftBorder")).set("color", "#225599")
    style = native.paragraph_styles()[0].element
    margin = style.find(qname("hh", "margin"))
    margin.find(qname("hc", "next")).set("value", "1200")
    switch = etree.SubElement(style, qname("hp", "switch"))
    case = etree.SubElement(switch, qname("hp", "case"))
    case.set(qname("hp", "required-namespace"), "http://www.hancom.co.kr/hwpml/2016/HwpUnitChar")
    case.append(deepcopy(margin))
    style.remove(margin)
    native.header.mark_modified()
    first = native.save(tmp_path / "source.hwpx")
    converted = Document.open(first).save(tmp_path / "middle.docx")
    result = HwpxDocument.open(Document.open(converted).save(tmp_path / "result.hwpx"))
    result.strict_validate()
    final = Document.open(result.save(tmp_path / "final.hwpx")).save(tmp_path / "final.docx")
    assert WordFile(final).paragraphs[0].paragraph_format.space_after == Pt(12)
    cell = WordFile(final).tables[0].cell(0, 0)
    assert cell._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:left")).get(qn("w:color")) == "225599"


def test_simple_equation_can_return_to_word(tmp_path: Path) -> None:
    word = WordFile()
    paragraph = word.add_paragraph("Equation ")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x+y=3"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    result, _ = _word_cycle(word, tmp_path)
    assert result.element.xpath(".//m:oMath/m:r/m:t/text()") == ["x+y=3"]


def test_fraction_scripts_and_radicals_are_not_flattened(tmp_path: Path) -> None:
    from jakal_hwpx._equation_conversion import hwp_to_omml, omml_to_hwp

    word = WordFile()
    paragraph = word.add_paragraph("Math ")
    script = "{a} over {b}+{c} over {d}+sqrt {{x} ^ {2}}+{y} _ {i} ^ {3}"
    original = hwp_to_omml(script)
    assert len(original.xpath("./m:f", namespaces={"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"})) == 2
    paragraph._p.append(original)
    result, _ = _word_cycle(word, tmp_path)
    math = result.element.xpath(".//m:oMath")[0]
    assert omml_to_hwp(math) == omml_to_hwp(original)
    namespaces = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    assert len(math.xpath("./m:f", namespaces=namespaces)) == 2
    assert math.xpath("./m:rad/m:e/m:sSup", namespaces=namespaces)
    assert math.xpath("./m:sSubSup/m:sub/m:r/m:t/text()", namespaces=namespaces) == ["i"]


def test_unsupported_word_math_fails_conversion_without_altering_original(tmp_path: Path) -> None:
    word = WordFile()
    math = OxmlElement("m:oMath")
    math.append(OxmlElement("m:nary"))
    word.add_paragraph()._p.append(math)
    source = tmp_path / "math.docx"
    word.save(source)
    document = Document.open(source)
    assert document.save(tmp_path / "copy.docx").read_bytes() == source.read_bytes()
    with pytest.raises(ValueError, match="Unsupported Word equation structure"):
        document.save(tmp_path / "unsupported.hwpx")


def test_picture_only_paragraph_keeps_alignment_in_both_directions(tmp_path: Path) -> None:
    word = WordFile()
    paragraph = word.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(BytesIO(PNG), width=Inches(1))
    result, _ = _word_cycle(word, tmp_path)
    assert len(result.paragraphs) == 1
    assert result.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(result.inline_shapes) == 1


def test_named_paragraph_styles_survive_both_directions(tmp_path: Path) -> None:
    word = WordFile()
    word.add_paragraph("Heading", style="Heading 1")
    word.add_paragraph("Body", style="Normal")
    result, native = _word_cycle(word, tmp_path)
    assert [p.style.name for p in result.paragraphs] == ["Heading 1", "Normal"]
    final = Document.open(tmp_path / "result.docx").save(tmp_path / "final.hwpx")
    assert [style.name for style in HwpxDocument.open(final).styles() if style.name in {"Heading 1", "Normal"}]


def test_custom_tab_stops_survive_three_roundtrips(tmp_path: Path) -> None:
    word = WordFile()
    paragraph = word.add_paragraph("Left\tRight")
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    for index in range(3):
        directory = tmp_path / str(index)
        directory.mkdir()
        word, _ = _word_cycle(word, directory)
        tabs = list(word.paragraphs[0].paragraph_format.tab_stops)
        assert len(tabs) == 1
        assert tabs[0].position == Inches(3)
        assert tabs[0].alignment == WD_TAB_ALIGNMENT.RIGHT
        assert tabs[0].leader == WD_TAB_LEADER.DOTS


def test_hyperlink_text_formatting_survives_roundtrip(tmp_path: Path) -> None:
    word = WordFile()
    paragraph = word.add_paragraph("Before ")
    _link(paragraph, "Styled link")
    run = paragraph.hyperlinks[0].runs[0]
    run.bold = True
    run.font.color.rgb = RGBColor(0x22, 0x55, 0x99)
    word, _ = _word_cycle(word, tmp_path)
    result = word.paragraphs[0].hyperlinks[0].runs[0]
    assert result.bold is True
    assert result.font.color.rgb == RGBColor(0x22, 0x55, 0x99)


def test_implicit_word_table_layout_does_not_become_centered_or_justified(tmp_path: Path) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

    word = WordFile()
    table = word.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "short line\nnext line"
    result, _ = _word_cycle(word, tmp_path)
    cell = result.tables[0].cell(0, 0)
    assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert cell.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.TOP
    assert cell._tc.tcPr.find(qn("w:tcMar")).find(qn("w:left")).get(qn("w:w")) == "108"


def test_metadata_and_empty_author_are_not_replaced_by_template_defaults(tmp_path: Path) -> None:
    word = WordFile()
    word.add_paragraph("Metadata")
    properties = word.core_properties
    properties.author = ""
    properties.title = "Title"
    properties.language = "ko-KR"
    properties.created = datetime(2020, 5, 3, 1, 2, 3, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 10, 28, 0, 0, 0, tzinfo=timezone.utc)
    result, _ = _word_cycle(word, tmp_path)
    assert result.core_properties.author == ""
    assert result.core_properties.title == "Title"
    assert result.core_properties.language == "ko-KR"
    assert result.core_properties.created == properties.created
    assert result.core_properties.modified == properties.modified
