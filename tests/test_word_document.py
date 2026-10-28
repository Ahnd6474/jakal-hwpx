from __future__ import annotations

from pathlib import Path
import io
import zipfile

from docx import Document as WordFile
from docx.oxml import OxmlElement
from lxml import etree
import pytest

from jakal_hwpx import Document, HancomDocument, WordDocument, HwpxDocument, HwpDocument, Paragraph

pytestmark = pytest.mark.filterwarnings("ignore:DOCX:UserWarning")


def test_document_factory_dispatches_to_word_subclass_and_keeps_alias(tmp_path: Path) -> None:
    source = WordFile()
    source.add_paragraph("Original")
    path = tmp_path / "source.docx"
    source.save(path)
    model = Document.open(path)
    assert HancomDocument is Document
    assert isinstance(model, Document)
    assert type(model) is WordDocument
    assert model.word_document.paragraphs[0].text == "Original"
    assert Document.blank().__class__ is Document
    hwpx = HwpxDocument.blank().save(tmp_path / "source.hwpx")
    assert type(Document.open(hwpx)) is Document


def test_shared_append_methods_modify_word_backend(tmp_path: Path) -> None:
    document = WordDocument.blank()
    document.append_paragraph("Shared paragraph")
    document.append_table(rows=1, cols=2, cell_texts=[["A", "B"]])
    document.append_hyperlink("https://example.com", display_text="Example")
    assert document.word_document.paragraphs[0].text == "Shared paragraph"
    assert document.word_document.tables[0].cell(0, 1).text == "B"
    output = document.save(tmp_path / "word.docx")
    assert WordFile(output).paragraphs[-1].text == "Example"
    document.append_header("Header")
    document.append_footer("Footer")
    assert document.word_document.sections[0].header.paragraphs[0].text == "Header"
    assert document.word_document.sections[0].footer.paragraphs[0].text == "Footer"


def test_word_native_edits_keep_unmodelled_package_content(tmp_path: Path) -> None:
    document = WordDocument.blank()
    document.append_paragraph("Before")
    source = document.save(tmp_path / "input.docx")
    with zipfile.ZipFile(source, "a") as package:
        package.writestr("customXml/opaque.xml", b"<kept/>")
    document = Document.open(source)
    with document.edit_word() as word:
        word.paragraphs[0].runs[0].bold = True
        word.add_paragraph("Native paragraph")
    assert any(style.bold for style in document.character_styles)
    assert any(isinstance(block, Paragraph) and block.text == "Native paragraph" for block in document.sections[0].blocks)
    result = document.save(tmp_path / "native.docx")
    with zipfile.ZipFile(result) as package:
        assert package.read("customXml/opaque.xml") == b"<kept/>"
    assert WordFile(result).paragraphs[0].runs[0].bold is True


def test_word_replacement_spans_runs_without_losing_rpr(tmp_path: Path) -> None:
    source = WordFile()
    paragraph = source.add_paragraph("Normal ")
    paragraph.add_run("bold").bold = True
    paragraph.add_run(" tail")
    path = tmp_path / "input.docx"
    source.save(path)
    model = Document.open(path)
    assert model.replace_text("Normal bold", "Edited bold") == 1
    assert model.word_document.paragraphs[0].text == "Edited bold tail"
    assert any(run.bold for run in model.word_document.paragraphs[0].runs)
    assert model.save(tmp_path / "edited.docx").is_file()


def test_word_conversion_uses_pure_python_by_default(tmp_path: Path) -> None:
    source = WordFile()
    source.add_paragraph("Word content")
    source.add_table(rows=1, cols=2).cell(0, 1).text = "Cell content"
    path = tmp_path / "source.docx"
    source.save(path)

    model = Document.open(path)
    hwpx = model.save(tmp_path / "word.hwpx")
    hwp = model.save(tmp_path / "word.hwp")
    assert "Word content" in HwpxDocument.open(hwpx).get_document_text()
    assert HwpxDocument.open(hwpx).tables()[0].cell(0, 1).text == "Cell content"
    assert "Word content" in HwpDocument.open(hwp).get_document_text()
    assert hwp.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1")


def test_word_character_effects_survive_pure_hwp_conversion(tmp_path: Path) -> None:
    source = WordFile()
    paragraph = source.add_paragraph()
    run = paragraph.add_run("Styled")
    run.bold = True
    run.italic = True
    run.underline = True
    run.font.strike = True
    run.font.superscript = True
    path = tmp_path / "styled.docx"
    source.save(path)

    document = Document.open(path)
    output = document.save(tmp_path / "styled.hwp")
    char_id = int(next(block.char_pr_id for block in document.sections[0].blocks if isinstance(block, Paragraph)))
    fields = HwpDocument.open(output).docinfo_model().char_shape_records()[char_id].fields()
    attributes = fields["attributes"]
    assert attributes & (1 << 0)  # italic
    assert attributes & (1 << 1)  # bold
    assert (attributes >> 2) & 0b11 == 1  # underline
    assert (attributes >> 18) & 0b111  # strikeout
    assert attributes & (1 << 15)  # superscript


def test_word_conversion_uses_explicit_converter_and_can_project(tmp_path: Path) -> None:
    source = WordFile()
    source.add_paragraph("Word content")
    path = tmp_path / "source.docx"
    source.save(path)
    calls = []

    def hancom(source_path, output_path, output_format):
        calls.append(output_format)
        assert WordFile(source_path).paragraphs[0].text == "Word content"
        result = HwpxDocument.blank()
        result.append_paragraph("Converted by Hancom")
        return result.save(output_path)

    model: Document = Document.open(path, converter=hancom)
    output = model.save(tmp_path / "word.hwpx")
    assert calls == ["HWPX"]
    assert "Converted by Hancom" in HwpxDocument.open(output).get_document_text()
    projected = model.to_common_document()
    assert type(projected) is Document
    assert "Word content" in projected.to_hwpx_document().get_document_text()
    assert calls == ["HWPX"]


def test_word_preserve_formatting_requires_explicit_converter(tmp_path: Path) -> None:
    source = WordFile()
    source.add_paragraph("Word content")
    path = tmp_path / "source.docx"
    source.save(path)
    with pytest.raises(ValueError, match="explicit converter"):
        Document.open(path).save(tmp_path / "word.hwpx", preserve_formatting=True)


def test_word_edit_context_rolls_back_on_error(tmp_path: Path) -> None:
    document = WordDocument.blank()
    document.append_paragraph("Kept")
    original = document.save(tmp_path / "before.docx").read_bytes()
    with pytest.raises(RuntimeError, match="abort"):
        with document.edit_word() as word:
            word.paragraphs[0].text = "Discarded"
            raise RuntimeError("abort")
    assert document.save(tmp_path / "after.docx").read_bytes() == original


def test_word_details_keep_comments_and_math_when_shared_text_is_edited(tmp_path: Path) -> None:
    source = WordFile()
    paragraph = source.add_paragraph("Before")
    source.add_comment(paragraph.runs[0], text="Keep this comment", author="Reviewer")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x+y"
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    path = tmp_path / "details.docx"
    source.save(path)
    with zipfile.ZipFile(path) as package:
        comments = package.read("word/comments.xml")
    document = Document.open(path)
    assert document.replace_text("Before", "After") == 1
    document.append_paragraph("New paragraph")
    output = document.save(tmp_path / "edited.docx")
    with zipfile.ZipFile(output) as package:
        assert package.read("word/comments.xml") == comments
        root = etree.fromstring(package.read("word/document.xml"))
        assert root.xpath("count(.//m:oMath)", namespaces={"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}) == 1
        assert b"x+y" in package.read("word/document.xml")


def test_word_metadata_edit_creates_reachable_core_part_if_original_omits_it(tmp_path: Path) -> None:
    source = WordFile()
    source.add_paragraph("Minimal")
    memory = io.BytesIO()
    source.save(memory)
    path = tmp_path / "minimal.docx"
    with zipfile.ZipFile(memory) as package, zipfile.ZipFile(path, "w") as output:
        for name in package.namelist():
            if name == "docProps/core.xml":
                continue
            data = package.read(name)
            if name in {"_rels/.rels", "[Content_Types].xml"}:
                root = etree.fromstring(data)
                for child in list(root):
                    if child.get("Target") == "docProps/core.xml" or child.get("PartName") == "/docProps/core.xml":
                        root.remove(child)
                data = etree.tostring(root)
            output.writestr(name, data)
    document = Document.open(path)
    assert document.save(tmp_path / "unmodified.docx").read_bytes() == path.read_bytes()
    document.metadata.title = "New title"
    result = document.save(tmp_path / "metadata.docx")
    assert WordFile(result).core_properties.title == "New title"


def test_word_both_page_header_dispatches_to_both_native_stories() -> None:
    document = WordDocument.blank()
    with document.edit_word() as word:
        word.settings.odd_and_even_pages_header_footer = True
    document.append_header("Both", apply_page_type="BOTH")
    assert document.word_document.sections[0].header.paragraphs[0].text == "Both"
    assert document.word_document.sections[0].even_page_header.paragraphs[0].text == "Both"
